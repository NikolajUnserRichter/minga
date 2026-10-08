"""Gernot-Feedback vom 07./08.10.2026 — Paket 1 (Steuer und Rechnung).

Gemeinsame Testdatei der Abschnitte S1–S6. Helfer und Klassen tragen ein
Abschnitts-Präfix (_s1_/TestS1, _s2_/TestS2, _s3_/TestS3, _datev_/TestDatev,
_s5_/TestS5, _s6_/TestS6): ein gleichnamiger Helfer würde still ersetzt.
"""
import csv
import io
import re
import uuid
from datetime import date, timedelta
from decimal import Decimal
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

import pytest

from tests.conftest import TestingSessionLocal
from tests.test_documents_preise import _pdf_text


# ============================================================
# S1 — Steuersatz kommt aus dem Produkt (A3)
#
# Pfandkisten liefen mit 7 % statt 19 %: das Bestellformular schickte fest
# REDUZIERT, das Schema setzte REDUZIERT als Default, die Rechnung aus der
# Bestellung übernahm den Satz der Bestellposition.
# ============================================================

_S1_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _s1_einheit():
    """Basiseinheit 'G' — Produktanlage und Produktimport verlangen sie."""
    from app.models.unit import UnitOfMeasure, UnitCategory
    with TestingSessionLocal() as db:
        unit = db.query(UnitOfMeasure).filter_by(code="G").first()
        if unit is None:
            unit = UnitOfMeasure(code="G", name="Gramm", symbol="g",
                                 category=UnitCategory.WEIGHT, is_base_unit=True)
            db.add(unit)
            db.commit()
        return str(unit.id)


def _s1_produkt(client, sku, name, preis, **extra):
    r = client.post("/api/v1/products", json={
        "sku": sku, "name": name, "base_price": str(preis),
        "category": "MICROGREEN", "base_unit_id": _s1_einheit(), **extra,
    })
    assert r.status_code in (200, 201), r.text
    return r.json()


def _s1_pfandkiste(client, sku="PFAND-IFCO", **extra):
    """Pfandkiste wie in Gernots Stamm: Kategorie PFAND → 19 % + Pfandkennzeichen."""
    return _s1_produkt(client, sku, "IFCO-Kiste", "3.00",
                       category="PFAND", deposit_value="3.00", **extra)


def _s1_variante(client, produkt):
    r = client.post(f"/api/v1/products/{produkt['id']}/variants", json={
        "name_suffix": "6er Stapel", "packaging_unit_id": produkt["base_unit_id"],
    })
    assert r.status_code in (200, 201), r.text
    return r.json()


def _s1_kunde(client, name="Ökoring Test", **extra):
    r = client.post("/api/v1/sales/customers", json={"name": name, "typ": "HANDEL", **extra})
    assert r.status_code in (200, 201), r.text
    return r.json()


def _s1_zeile(produkt=None, menge=2, preis="3.00", **extra):
    zeile = {
        "product_name": produkt["name"] if produkt else "Kresse (Freitext)",
        "quantity": menge, "unit": "STK", "unit_price": preis,
    }
    if produkt:
        zeile["product_id"] = produkt["id"]
    zeile.update(extra)
    return zeile


def _s1_bestellung(client, kunde, zeilen, liefertag="2026-10-09"):
    r = client.post("/api/v1/sales/orders", json={
        "customer_id": kunde["id"], "requested_delivery_date": liefertag, "lines": zeilen,
    })
    assert r.status_code == 201, r.text
    return r.json()


def _s1_altstand(order_id):
    """Stellt den Stand vor dem Fix her: jede Produktposition auf 7 %, Beträge
    und Summen passend dazu — so sehen Gernots offene Bestellungen aus."""
    from app.api.v1.sales import _calculate_line_amounts, _calculate_order_totals
    from app.models.enums import TaxRate
    from app.models.order import Order
    with TestingSessionLocal() as db:
        order = db.get(Order, uuid.UUID(order_id))
        for line in order.lines:
            if line.product_id or line.product_variant_id:
                line.tax_rate = TaxRate.REDUZIERT
                _calculate_line_amounts(line)
        _calculate_order_totals(order)
        db.commit()


def _s1_altbestellung(client, kunde, pfand, liefertag="2026-10-09"):
    """Offene Bestellung wie vor dem Fix: Kresse 7 %, Pfandkiste fälschlich 7 %."""
    order = _s1_bestellung(client, kunde, [
        _s1_zeile(None, menge=10, preis="2.50", tax_rate="REDUZIERT"),
        _s1_zeile(pfand, menge=2, preis="3.00"),
    ], liefertag=liefertag)
    _s1_altstand(order["id"])
    return client.get(f"/api/v1/sales/orders/{order['id']}").json()


def _s1_pfandzeile(client, order_id):
    order = client.get(f"/api/v1/sales/orders/{order_id}").json()
    return next(z for z in order["lines"] if z["product_id"])


def _s1_rechnung_aus(client, order):
    r = client.post(f"/api/v1/invoices/from-order/{order['id']}")
    assert r.status_code in (200, 201), r.text
    return r.json()


def _s1_rechnungszeilen(invoice_id):
    """Rechnungszeilen direkt aus der DB — is_deposit steht nicht in der API-Antwort."""
    from app.models.invoice import InvoiceLine
    with TestingSessionLocal() as db:
        return [{
            "product_id": l.product_id,
            "tax_rate": l.tax_rate.value,
            "is_deposit": l.is_deposit,
            "discount_percent": Decimal(str(l.discount_percent)),
        } for l in db.query(InvoiceLine)
            .filter_by(invoice_id=uuid.UUID(invoice_id))
            .order_by(InvoiceLine.position)]


def _s1_produktimport(client, *zeilen):
    """Produktimport mit der Kopfzeile des echten Templates; je Zeile nur die
    angegebenen Spalten befüllt, alle anderen Zellen leer."""
    from openpyxl import Workbook, load_workbook
    tpl = client.get("/api/v1/imports/template/products").content
    header = [c.value for c in load_workbook(io.BytesIO(tpl))["Daten"][1]]
    spalten = [str(h).rstrip(" *") for h in header]
    wb = Workbook()
    ws = wb.active
    ws.title = "Daten"
    ws.append(header)
    for werte in zeilen:
        zeile = [""] * len(header)
        for spalte, wert in werte.items():
            zeile[spalten.index(spalte)] = wert
        ws.append(zeile)
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/api/v1/imports/products",
                    files={"file": ("products.xlsx", buf.getvalue(), _S1_XLSX)})
    assert r.status_code == 200, r.text
    return r.json()


def _s1_produkt_orm(db, sku, name, satz):
    """Produkt direkt über das ORM — für Service-Tests mit der db-Fixture ohne Client."""
    from app.models.product import Product, ProductCategory
    from app.models.unit import UnitOfMeasure, UnitCategory
    unit = db.query(UnitOfMeasure).filter_by(code="STK").first()
    if unit is None:
        unit = UnitOfMeasure(code="STK", name="Stück", category=UnitCategory.COUNT)
        db.add(unit)
        db.flush()
    produkt = Product(sku=sku, name=name, category=ProductCategory.MICROGREEN,
                      base_unit_id=unit.id, tax_rate=satz)
    db.add(produkt)
    db.commit()
    return produkt


class TestS1SteuersatzInDerBestellung:
    """Task 1: Bestellpositionen mit Produkt tragen den Produktsatz."""

    def test_pfandkiste_bekommt_19_prozent_obwohl_das_formular_7_schickt(self, client):
        pfand = _s1_pfandkiste(client)
        order = _s1_bestellung(client, _s1_kunde(client), [
            _s1_zeile(pfand, menge=2, preis="3.00", tax_rate="REDUZIERT"),
        ])
        zeile = order["lines"][0]
        assert zeile["tax_rate"] == "STANDARD"
        # 2 × 3,00 € = 6,00 € netto, 19 % = 1,14 €
        assert Decimal(str(zeile["line_vat"])) == Decimal("1.14")
        assert Decimal(str(order["total_gross"])) == Decimal("7.14")

    def test_ohne_satz_gilt_der_produktsatz(self, client):
        pfand = _s1_pfandkiste(client)
        order = _s1_bestellung(client, _s1_kunde(client), [_s1_zeile(pfand)])
        assert order["lines"][0]["tax_rate"] == "STANDARD"

    def test_bewusst_auf_7_prozent_gesetztes_pfand_bleibt_bei_7(self, client):
        """Der Fix übernimmt den Produktsatz — er erzwingt nicht 19 % für Pfand."""
        pfand = _s1_pfandkiste(client, sku="PFAND-7", tax_rate="REDUZIERT")
        order = _s1_bestellung(client, _s1_kunde(client), [_s1_zeile(pfand, tax_rate="STANDARD")])
        assert order["lines"][0]["tax_rate"] == "REDUZIERT"

    def test_freitext_behaelt_den_satz_vom_client(self, client):
        order = _s1_bestellung(client, _s1_kunde(client), [
            _s1_zeile(None, tax_rate="STANDARD"),
            {"product_name": "Freitext ohne Satz", "quantity": 1, "unit": "STK", "unit_price": "2.00"},
        ])
        assert sorted(z["tax_rate"] for z in order["lines"]) == ["REDUZIERT", "STANDARD"]

    def test_position_nur_mit_variante_nimmt_den_satz_des_elternprodukts(self, client):
        pfand = _s1_pfandkiste(client)
        variante = _s1_variante(client, pfand)
        order = _s1_bestellung(client, _s1_kunde(client), [{
            "product_variant_id": variante["id"], "product_name": "IFCO-Kiste",
            "quantity": 1, "unit": "STK", "unit_price": "18.00", "tax_rate": "REDUZIERT",
        }])
        assert order["lines"][0]["tax_rate"] == "STANDARD"

    def test_nachgetragene_position_nimmt_den_produktsatz(self, client):
        """EditOrderModal schickt keinen Satz — bisher griff der Schema-Default 7 %."""
        pfand = _s1_pfandkiste(client)
        order = _s1_bestellung(client, _s1_kunde(client), [
            _s1_zeile(None, menge=10, preis="2.50", tax_rate="REDUZIERT"),
        ])
        r = client.post(f"/api/v1/sales/orders/{order['id']}/lines", json={
            "product_id": pfand["id"], "product_name": pfand["name"],
            "quantity": 2, "unit": "STK", "unit_price": "3.00",
        })
        assert r.status_code == 201, r.text
        assert r.json()["tax_rate"] == "STANDARD"
        detail = client.get(f"/api/v1/sales/orders/{order['id']}").json()
        # 25,00 € × 7 % = 1,75 € | 6,00 € × 19 % = 1,14 €
        assert Decimal(str(detail["total_vat"])) == Decimal("2.89")

    def test_nachgetragene_variantenposition_behaelt_die_variante(self, client):
        """add_order_line verwarf product_variant_id — die Rechnung fand dann
        kein Produkt hinter der Position und setzte kein Pfandkennzeichen."""
        pfand = _s1_pfandkiste(client)
        variante = _s1_variante(client, pfand)
        order = _s1_bestellung(client, _s1_kunde(client), [
            _s1_zeile(None, menge=10, preis="2.50", tax_rate="REDUZIERT"),
        ])
        r = client.post(f"/api/v1/sales/orders/{order['id']}/lines", json={
            "product_variant_id": variante["id"], "product_name": "IFCO-Kiste",
            "quantity": 1, "unit": "STK", "unit_price": "18.00",
        })
        assert r.status_code == 201, r.text
        assert r.json()["product_variant_id"] == variante["id"]

        detail = client.get(f"/api/v1/sales/orders/{order['id']}").json()
        zeile = next(z for z in detail["lines"] if z["product_variant_id"])
        assert zeile["product_variant_id"] == variante["id"]
        assert zeile["tax_rate"] == "STANDARD"

        # Unbekannte Variante: 404 wie in create_order, kein Fremdschlüsselfehler
        r = client.post(f"/api/v1/sales/orders/{order['id']}/lines", json={
            "product_variant_id": str(uuid.uuid4()), "product_name": "?",
            "quantity": 1, "unit": "STK", "unit_price": "1.00",
        })
        assert r.status_code == 404, r.text

    def test_patch_kann_den_produktsatz_nicht_ueberschreiben(self, client):
        pfand = _s1_pfandkiste(client)
        order = _s1_bestellung(client, _s1_kunde(client), [_s1_zeile(pfand)])
        zeile = order["lines"][0]
        r = client.patch(f"/api/v1/sales/orders/{order['id']}/lines/{zeile['id']}",
                         json={"tax_rate": "REDUZIERT"})
        assert r.status_code == 200, r.text
        assert r.json()["tax_rate"] == "STANDARD"

    def test_patch_zieht_den_alten_satz_nach_und_protokolliert_ihn(self, client):
        """Bestätigte Bestellung mit Altposition (7 %): die nächste Änderung
        bringt den Produktsatz und steht mit altem und neuem Satz im Audit-Log."""
        pfand = _s1_pfandkiste(client)
        order = _s1_bestellung(client, _s1_kunde(client), [_s1_zeile(pfand)])
        zeile = order["lines"][0]
        _s1_altstand(order["id"])
        r = client.post(f"/api/v1/sales/orders/{order['id']}/confirm")
        assert r.status_code == 200, r.text

        r = client.patch(f"/api/v1/sales/orders/{order['id']}/lines/{zeile['id']}",
                         json={"quantity": 3})
        assert r.status_code == 200, r.text
        assert r.json()["tax_rate"] == "STANDARD"

        log = client.get(f"/api/v1/sales/orders/{order['id']}/audit-log").json()
        eintrag = next(e for e in log if e["action"] == "UPDATE_LINE")
        assert eintrag["old_values"]["tax_rate"] == "REDUZIERT"
        assert eintrag["new_values"]["tax_rate"] == "STANDARD"

    def test_kreditlimit_rechnet_pfand_mit_19_prozent(self, client):
        """100 € Pfand netto: mit 7 % 107 € (unter dem Limit von 110 €), richtig 119 €."""
        pfand = _s1_pfandkiste(client)
        kunde = _s1_kunde(client, credit_limit="110.00")
        r = client.post("/api/v1/sales/orders", json={
            "customer_id": kunde["id"], "requested_delivery_date": "2026-10-09",
            "lines": [_s1_zeile(pfand, menge=1, preis="100.00", tax_rate="REDUZIERT")],
        })
        assert r.status_code == 400, r.text
        assert "Kreditlimit" in r.json()["detail"]


class TestS1RechnungAusBestellung:
    """Task 2: Rechnung aus der Bestellung — Satz, Rabatt, Leistungsdatum."""

    def test_altposition_mit_7_prozent_wird_mit_dem_produktsatz_fakturiert(self, client):
        """Der Weg von RE-00002/3/4: Bestellposition 7 %, Produkt 19 %."""
        pfand = _s1_pfandkiste(client)
        order = _s1_altbestellung(client, _s1_kunde(client), pfand)

        rechnung = _s1_rechnung_aus(client, order)

        zeilen = _s1_rechnungszeilen(rechnung["id"])
        pfandzeile = next(z for z in zeilen if z["product_id"] is not None)
        assert pfandzeile["tax_rate"] == "STANDARD"
        assert pfandzeile["is_deposit"] is True
        assert next(z for z in zeilen if z["product_id"] is None)["tax_rate"] == "REDUZIERT"
        # 25,00 € × 7 % = 1,75 € | 6,00 € × 19 % = 1,14 €
        assert Decimal(str(rechnung["tax_amount"])) == Decimal("2.89")

    def test_variantenposition_traegt_satz_und_pfandkennzeichen(self, client):
        pfand = _s1_pfandkiste(client)
        variante = _s1_variante(client, pfand)
        order = _s1_bestellung(client, _s1_kunde(client), [{
            "product_variant_id": variante["id"], "product_name": "IFCO-Kiste",
            "quantity": 1, "unit": "STK", "unit_price": "18.00",
        }])
        _s1_altstand(order["id"])

        zeile = _s1_rechnungszeilen(_s1_rechnung_aus(client, order)["id"])[0]

        assert zeile["tax_rate"] == "STANDARD"
        assert zeile["is_deposit"] is True
        assert zeile["product_id"] == uuid.UUID(pfand["id"])

    def test_positionsrabatt_wird_uebernommen(self, client):
        order = _s1_bestellung(client, _s1_kunde(client), [
            _s1_zeile(None, menge=10, preis="10.00", tax_rate="REDUZIERT", discount_percent="10"),
        ])
        assert Decimal(str(order["total_net"])) == Decimal("90.00")

        rechnung = _s1_rechnung_aus(client, order)

        assert _s1_rechnungszeilen(rechnung["id"])[0]["discount_percent"] == Decimal("10")
        assert Decimal(str(rechnung["subtotal"])) == Decimal("90.00")

    def test_leistungsdatum_ist_das_tatsaechliche_lieferdatum(self, client):
        from app.models.order import Order
        order = _s1_bestellung(client, _s1_kunde(client), [_s1_zeile(None, tax_rate="REDUZIERT")],
                               liefertag="2026-10-07")
        with TestingSessionLocal() as db:
            db.get(Order, uuid.UUID(order["id"])).actual_delivery_date = date(2026, 10, 8)
            db.commit()

        assert _s1_rechnung_aus(client, order)["delivery_date"] == "2026-10-08"

    def test_ohne_lieferung_bleibt_das_wunschlieferdatum(self, client):
        order = _s1_bestellung(client, _s1_kunde(client), [_s1_zeile(None, tax_rate="REDUZIERT")],
                               liefertag="2026-10-07")
        assert _s1_rechnung_aus(client, order)["delivery_date"] == "2026-10-07"

    def test_storno_bucht_mit_dem_satz_der_originalzeile_gegen(self, client):
        """Absicherung: der Produktsatz wird NICHT in add_line erzwungen. Der
        Storno einer alten 7-%-Pfandrechnung muss mit 7 % gegenbuchen, sonst
        ergibt er nicht null."""
        pfand = _s1_pfandkiste(client)
        kunde = _s1_kunde(client)
        rechnung = client.post("/api/v1/invoices", json={
            "customer_id": kunde["id"], "invoice_date": "2026-10-08",
        }).json()
        r = client.post(f"/api/v1/invoices/{rechnung['id']}/lines", json={
            "description": "IFCO-Kiste", "quantity": 2, "unit": "STK", "unit_price": "3.00",
            "tax_rate": "REDUZIERT", "product_id": pfand["id"],
        })
        assert r.status_code == 201, r.text
        r = client.post(f"/api/v1/invoices/{rechnung['id']}/finalize")
        assert r.status_code == 200, r.text

        storno = client.post(f"/api/v1/invoices/{rechnung['id']}/cancel",
                             json={"reason": "Steuersatz falsch"})
        assert storno.status_code == 200, storno.text

        gutschrift = storno.json()["credit_note"]
        assert [z["tax_rate"] for z in _s1_rechnungszeilen(gutschrift["id"])] == ["REDUZIERT"]


class TestS1Sammelrechnung:
    """Task 3: Sammelrechnung übergibt Produkt, Produktsatz und Rabatt."""

    def _lauf(self, client, kunde, *zeilen):
        order = _s1_bestellung(client, kunde, list(zeilen), liefertag="2026-09-15")
        _s1_altstand(order["id"])
        note = client.post(f"/api/v1/sales/orders/{order['id']}/delivery-notes", json={})
        assert note.status_code == 201, note.text
        r = client.post("/api/v1/invoices/batch-run/commit", json={
            "period_from": "2026-09-01", "period_to": "2026-09-30",
            "customer_ids": [kunde["id"]],
        })
        assert r.status_code == 201, r.text
        return r.json()["rechnungen"][0]

    def test_pfandzeile_traegt_produkt_satz_und_kennzeichen(self, client):
        pfand = _s1_pfandkiste(client)
        rechnung = self._lauf(client, _s1_kunde(client), _s1_zeile(pfand, menge=4, preis="3.00"))

        zeile = _s1_rechnungszeilen(rechnung["id"])[0]
        assert zeile["product_id"] == uuid.UUID(pfand["id"])
        assert zeile["is_deposit"] is True
        assert zeile["tax_rate"] == "STANDARD"
        # 4 × 3,00 € netto + 19 % = 14,28 € Pfand brutto
        assert Decimal(str(rechnung["total_deposit"])) == Decimal("14.28")

    def test_positionsrabatt_bleibt_erhalten(self, client):
        rechnung = self._lauf(client, _s1_kunde(client), _s1_zeile(
            None, menge=10, preis="10.00", tax_rate="REDUZIERT", discount_percent="10"))
        assert Decimal(str(rechnung["subtotal"])) == Decimal("90.00")


class TestS1Produktstamm:
    """Task 4: PATCH und Import wenden dieselbe Pfandregel an wie das Anlegen."""

    def test_patch_auf_pfandkennzeichen_setzt_19_prozent(self, client):
        kiste = _s1_produkt(client, "KISTE-1", "Mehrwegkiste", "3.00", category="PACKAGING")
        assert kiste["tax_rate"] == "REDUZIERT"
        r = client.patch(f"/api/v1/products/{kiste['id']}", json={"is_deposit": True})
        assert r.status_code == 200, r.text
        assert r.json()["tax_rate"] == "STANDARD"

    def test_patch_auf_kategorie_pfand_setzt_kennzeichen_und_19_prozent(self, client):
        kiste = _s1_produkt(client, "KISTE-2", "Mehrwegkiste", "3.00", category="PACKAGING")
        r = client.patch(f"/api/v1/products/{kiste['id']}", json={"category": "PFAND"})
        assert r.status_code == 200, r.text
        assert r.json()["category"] == "PFAND"
        assert r.json()["is_deposit"] is True
        assert r.json()["tax_rate"] == "STANDARD"

    def test_ausdruecklicher_satz_im_patch_gewinnt(self, client):
        kiste = _s1_produkt(client, "KISTE-3", "Mehrwegkiste", "3.00", category="PACKAGING")
        r = client.patch(f"/api/v1/products/{kiste['id']}",
                         json={"is_deposit": True, "tax_rate": "REDUZIERT"})
        assert r.status_code == 200, r.text
        assert r.json()["tax_rate"] == "REDUZIERT"

    def test_patch_mit_leerer_kategorie_wird_abgewiesen(self, client):
        """products.category ist NOT NULL: ein ausdrückliches null ergibt 422,
        nicht einen Datenbankfehler (500) beim Commit."""
        kiste = _s1_produkt(client, "KISTE-4", "Mehrwegkiste", "3.00", category="PACKAGING")
        r = client.patch(f"/api/v1/products/{kiste['id']}", json={"category": None})
        assert r.status_code == 422, r.text
        assert client.get(f"/api/v1/products/{kiste['id']}").json()["category"] == "PACKAGING"

    def test_speichern_eines_7_prozent_pfands_kippt_den_satz_nicht(self, client):
        """Erneutes is_deposit=true ist kein Übergang — der bewusst gewählte Satz bleibt."""
        pfand = _s1_pfandkiste(client, sku="PFAND-7B", tax_rate="REDUZIERT")
        r = client.patch(f"/api/v1/products/{pfand['id']}",
                         json={"is_deposit": True, "name": "IFCO-Kiste 7 %"})
        assert r.status_code == 200, r.text
        assert r.json()["tax_rate"] == "REDUZIERT"

    def test_reimport_ohne_steuersatz_laesst_pfand_bei_19_prozent(self, client):
        pfand = _s1_pfandkiste(client, sku="PFAND-IMP1")
        assert pfand["tax_rate"] == "STANDARD"

        ergebnis = _s1_produktimport(client, {
            "sku": "PFAND-IMP1", "name": "IFCO-Kiste", "category": "PFAND", "base_price": "3.10",
        })

        assert ergebnis["updated"] == 1, ergebnis
        p = client.get(f"/api/v1/products/{pfand['id']}").json()
        assert p["tax_rate"] == "STANDARD"
        assert Decimal(str(p["base_price"])) == Decimal("3.10")

    def test_neuer_pfandartikel_ohne_steuersatz_bekommt_19_prozent(self, client):
        _s1_einheit()
        ergebnis = _s1_produktimport(client, {"sku": "PFAND-IMP2", "name": "IFCO-Kiste", "category": "PFAND"})
        assert ergebnis["created"] == 1, ergebnis
        p = next(p for p in client.get("/api/v1/products").json() if p["sku"] == "PFAND-IMP2")
        assert p["is_deposit"] is True
        assert p["tax_rate"] == "STANDARD"

    def test_zeile_ohne_kategorie_laesst_den_bestand_unveraendert(self, client):
        """Gegenprobe zur Lückenprüfung: category ist Pflichtspalte, eine Zeile
        ohne sie wird verworfen — der Pfandartikel bleibt PFAND."""
        pfand = _s1_pfandkiste(client, sku="PFAND-IMP3")
        ergebnis = _s1_produktimport(client, {"sku": "PFAND-IMP3", "name": "IFCO-Kiste"})
        assert ergebnis["updated"] == 0
        assert any("category" in f for f in ergebnis["errors"]), ergebnis
        p = client.get(f"/api/v1/products/{pfand['id']}").json()
        assert p["category"] == "PFAND"
        assert p["tax_rate"] == "STANDARD"


class TestS1Shopify:
    """Task 5: Shopify-Import nimmt den Produktsatz statt pauschal 19 %."""

    def _bestellung(self, *positionen):
        return {
            "id": 9001, "name": "#9001", "currency": "EUR",
            "customer": {"first_name": "Eva", "last_name": "Shop", "email": "eva@example.com"},
            "line_items": list(positionen),
        }

    def _saetze(self, db):
        from app.models.order import Order
        order = db.query(Order).filter_by(customer_reference="#9001").one()
        return [l.tax_rate.value for l in sorted(order.lines, key=lambda l: l.position)]

    def test_bekannter_artikel_nimmt_den_produktsatz(self, db):
        from app.models.enums import TaxRate
        from app.services.shopify_service import import_shopify_order
        _s1_produkt_orm(db, "MG-ERBSE", "Erbsenkresse", TaxRate.REDUZIERT)

        import_shopify_order(db, self._bestellung(
            {"title": "Erbsenkresse", "sku": "MG-ERBSE", "quantity": 2, "price": "4.00"}))

        assert self._saetze(db) == ["REDUZIERT"]

    def test_unbekannter_artikel_nimmt_den_satz_aus_shopify(self, db):
        from app.services.shopify_service import import_shopify_order
        import_shopify_order(db, self._bestellung(
            {"title": "Kresse lose", "sku": None, "quantity": 1, "price": "3.00",
             "tax_lines": [{"rate": 0.07, "price": "0.21", "title": "MwSt"}]},
            {"title": "Tasse", "sku": None, "quantity": 1, "price": "9.00"}))

        # ohne tax_lines wie bisher 19 %
        assert self._saetze(db) == ["REDUZIERT", "STANDARD"]


class TestS1Datenkorrektur:
    """Task 6: einmalige Korrektur offener Bestellpositionen."""

    def _lauf(self, **kw):
        from app.services.steuersatz_korrektur import korrigiere_offene_bestellpositionen
        with TestingSessionLocal() as db:
            ergebnis = korrigiere_offene_bestellpositionen(db, **kw)
            db.commit()
            return ergebnis

    def test_offene_bestellung_wird_auf_den_produktsatz_korrigiert(self, client):
        pfand = _s1_pfandkiste(client)
        order = _s1_altbestellung(client, _s1_kunde(client), pfand)
        # 31,00 € × 7 % = 2,17 € — der falsche Stand
        assert Decimal(str(order["total_vat"])) == Decimal("2.17")

        ergebnis = self._lauf()

        assert ergebnis["bestellungen"] == 1 and ergebnis["positionen"] == 1
        nachher = client.get(f"/api/v1/sales/orders/{order['id']}").json()
        assert {bool(z["product_id"]): z["tax_rate"] for z in nachher["lines"]} == {
            True: "STANDARD", False: "REDUZIERT",
        }
        pfandzeile = _s1_pfandzeile(client, order["id"])
        assert Decimal(str(pfandzeile["line_vat"])) == Decimal("1.14")
        assert Decimal(str(pfandzeile["line_gross"])) == Decimal("7.14")
        # 25,00 € × 7 % = 1,75 € | 6,00 € × 19 % = 1,14 €
        assert Decimal(str(nachher["total_vat"])) == Decimal("2.89")
        assert Decimal(str(nachher["total_gross"])) == Decimal("33.89")

    def test_korrektur_steht_im_audit_log(self, client):
        pfand = _s1_pfandkiste(client)
        order = _s1_altbestellung(client, _s1_kunde(client), pfand)

        self._lauf()

        log = client.get(f"/api/v1/sales/orders/{order['id']}/audit-log").json()
        eintrag = next(e for e in log if e["action"] == "STEUERSATZ_KORREKTUR")
        assert eintrag["old_values"]["positionen"][0]["tax_rate"] == "REDUZIERT"
        assert eintrag["new_values"]["positionen"][0]["tax_rate"] == "STANDARD"
        assert Decimal(eintrag["old_values"]["total_gross"]) == Decimal("33.17")
        assert Decimal(eintrag["new_values"]["total_gross"]) == Decimal("33.89")

    def test_zweiter_lauf_aendert_nichts(self, client):
        pfand = _s1_pfandkiste(client)
        order = _s1_altbestellung(client, _s1_kunde(client), pfand)
        self._lauf()

        zweiter = self._lauf()

        assert zweiter["bestellungen"] == 0
        log = client.get(f"/api/v1/sales/orders/{order['id']}/audit-log").json()
        assert sum(1 for e in log if e["action"] == "STEUERSATZ_KORREKTUR") == 1

    def test_bestellung_mit_rechnung_wird_nicht_angefasst(self, client):
        """Auch ein Rechnungsentwurf (RE-00003) sperrt: Rechnungen fasst die
        Korrektur nie an, und die Bestellung soll nicht von ihr abweichen."""
        pfand = _s1_pfandkiste(client)
        order = _s1_altbestellung(client, _s1_kunde(client), pfand)
        rechnung = _s1_rechnung_aus(client, order)
        rechnungszeilen = _s1_rechnungszeilen(rechnung["id"])

        ergebnis = self._lauf()

        assert ergebnis["bestellungen"] == 0
        assert ergebnis["uebersprungen_mit_rechnung"] == [order["order_number"]]
        assert _s1_pfandzeile(client, order["id"])["tax_rate"] == "REDUZIERT"
        assert _s1_rechnungszeilen(rechnung["id"]) == rechnungszeilen

    def test_bestellung_in_sammelrechnung_wird_nicht_angefasst(self, client):
        pfand = _s1_pfandkiste(client)
        kunde = _s1_kunde(client)
        order = _s1_altbestellung(client, kunde, pfand, liefertag="2026-09-15")
        assert client.post(f"/api/v1/sales/orders/{order['id']}/delivery-notes", json={}).status_code == 201
        r = client.post("/api/v1/invoices/batch-run/commit", json={
            "period_from": "2026-09-01", "period_to": "2026-09-30", "customer_ids": [kunde["id"]],
        })
        assert r.status_code == 201, r.text

        ergebnis = self._lauf()

        assert ergebnis["bestellungen"] == 0
        assert _s1_pfandzeile(client, order["id"])["tax_rate"] == "REDUZIERT"

    def test_stornierte_bestellung_bleibt(self, client):
        from app.models.enums import OrderStatus
        from app.models.order import Order
        pfand = _s1_pfandkiste(client)
        order = _s1_altbestellung(client, _s1_kunde(client), pfand)
        with TestingSessionLocal() as db:
            db.get(Order, uuid.UUID(order["id"])).status = OrderStatus.STORNIERT
            db.commit()

        assert self._lauf()["bestellungen"] == 0
        assert _s1_pfandzeile(client, order["id"])["tax_rate"] == "REDUZIERT"

    def test_nach_storno_der_rechnung_gezielt_korrigierbar(self, client):
        """Für die Neuausstellung von RE-00002/4: erst stornieren, dann die
        Bestellung gezielt korrigieren, dann neu fakturieren."""
        pfand = _s1_pfandkiste(client)
        order = _s1_altbestellung(client, _s1_kunde(client), pfand)
        rechnung = _s1_rechnung_aus(client, order)
        r = client.post(f"/api/v1/invoices/{rechnung['id']}/finalize")
        assert r.status_code == 200, r.text
        r = client.post(f"/api/v1/invoices/{rechnung['id']}/cancel", json={"reason": "Steuersatz Pfand"})
        assert r.status_code == 200, r.text

        ergebnis = self._lauf(nur_bestellungen=[uuid.UUID(order["id"])])

        assert ergebnis["bestellungen"] == 1
        assert ergebnis["geaendert"] == [order["order_number"]]
        # Die Stornorechnung sperrt nicht: beide Listen leer, wie das Runbook es verlangt
        assert ergebnis["uebersprungen_mit_rechnung"] == []
        assert ergebnis["uebersprungen_status"] == []
        assert _s1_pfandzeile(client, order["id"])["tax_rate"] == "STANDARD"

    def test_gezielter_aufruf_meldet_nicht_offene_und_unbekannte_bestellungen(self, client):
        """nur_bestellungen filtert nach Status — eine fakturierte Bestellung
        oder eine unbekannte ID darf nicht still durchrutschen. Sonst hielte
        das Runbook ein leeres Ergebnis für Erfolg."""
        from app.models.enums import OrderStatus
        from app.models.order import Order
        pfand = _s1_pfandkiste(client)
        order = _s1_altbestellung(client, _s1_kunde(client), pfand)
        with TestingSessionLocal() as db:
            db.get(Order, uuid.UUID(order["id"])).status = OrderStatus.FAKTURIERT
            db.commit()
        unbekannt = uuid.uuid4()

        ergebnis = self._lauf(nur_bestellungen=[uuid.UUID(order["id"]), unbekannt])

        assert ergebnis["bestellungen"] == 0
        assert ergebnis["uebersprungen_status"] == [
            f"{order['order_number']} (FAKTURIERT)",
            f"{unbekannt} (nicht gefunden)",
        ]
        assert _s1_pfandzeile(client, order["id"])["tax_rate"] == "REDUZIERT"
        # Der automatische Lauf meldet fakturierte Bestellungen nicht.
        assert self._lauf()["uebersprungen_status"] == []

    def test_laeuft_je_mandant_nur_einmal(self, client):
        from app.models.app_setting import AppSetting
        from app.services.steuersatz_korrektur import MARKER, korrektur_einmalig_ausfuehren
        pfand = _s1_pfandkiste(client)
        order = _s1_altbestellung(client, _s1_kunde(client), pfand)
        with TestingSessionLocal() as db:
            erstes = korrektur_einmalig_ausfuehren(db)
            db.commit()
        assert erstes["bestellungen"] == 1

        _s1_altstand(order["id"])  # dieselbe Abweichung taucht wieder auf
        with TestingSessionLocal() as db:
            zweites = korrektur_einmalig_ausfuehren(db)
            db.commit()
            assert db.get(AppSetting, MARKER) is not None

        assert zweites is None
        assert _s1_pfandzeile(client, order["id"])["tax_rate"] == "REDUZIERT"

    def test_auto_migrate_fuehrt_die_korrektur_aus(self):
        """Verdrahtung in tenancy._auto_migrate — auf einer eigenen In-Memory-
        Engine, die geteilte Test-Engine bleibt unberührt."""
        from sqlalchemy import create_engine
        from sqlalchemy.orm import Session
        from sqlalchemy.pool import StaticPool
        from app.database import Base
        from app.models.app_setting import AppSetting
        from app.services.steuersatz_korrektur import MARKER
        from app.tenancy import _auto_migrate

        engine = create_engine("sqlite://", connect_args={"check_same_thread": False},
                               poolclass=StaticPool)
        Base.metadata.create_all(bind=engine)
        try:
            _auto_migrate(engine)
            with Session(engine) as s:
                assert s.get(AppSetting, MARKER) is not None
        finally:
            Base.metadata.drop_all(bind=engine)
            engine.dispose()


    def test_gescheiterte_korrektur_blockiert_keine_schemamigration(self, monkeypatch, caplog):
        """Scheitert die Korrektur, laufen die Spaltenmigrationen trotzdem, es
        steht kein Marker, und _auto_migrate wirft nicht — beim nächsten Start
        wird sie erneut versucht. Stünde der Block am Anfang des ersten try,
        fehlte die Spalte; stünde er in dessen except-Zweig, fehlte die
        eigene Fehlermeldung."""
        import logging
        from sqlalchemy import create_engine, inspect, text
        from sqlalchemy.orm import Session
        from sqlalchemy.pool import StaticPool
        from app.database import Base
        from app.models.app_setting import AppSetting
        from app.services.steuersatz_korrektur import MARKER
        from app.tenancy import _auto_migrate

        def scheitert(db, nur_bestellungen=None):
            raise RuntimeError("Korrektur gescheitert (Test)")

        monkeypatch.setattr(
            "app.services.steuersatz_korrektur.korrigiere_offene_bestellpositionen", scheitert)

        engine = create_engine("sqlite://", connect_args={"check_same_thread": False},
                               poolclass=StaticPool)
        Base.metadata.create_all(bind=engine)
        try:
            # Älteres Schema: eine Spalte fehlt, die _auto_migrate nachträgt.
            with engine.begin() as conn:
                conn.execute(text("ALTER TABLE invoices DROP COLUMN service_period_end"))

            with caplog.at_level(logging.ERROR, logger="app.tenancy"):
                _auto_migrate(engine)  # wirft nicht

            spalten = {c["name"] for c in inspect(engine).get_columns("invoices")}
            assert "service_period_end" in spalten
            with Session(engine) as s:
                assert s.get(AppSetting, MARKER) is None
            assert "Steuersatz-Korrektur fehlgeschlagen" in caplog.text
        finally:
            Base.metadata.drop_all(bind=engine)
            engine.dispose()


# ===========================================================================
# S2 — Rechnungs-PDF weist Steuer je Satz aus, Rundung einheitlich
# Helfer und Klassen tragen das Präfix _s2_/TestS2, weil sich mehrere
# Abschnitte diese Datei teilen: ein gleichnamiger Helfer oder eine
# gleichnamige Klasse aus einem anderen Abschnitt würde still ersetzt.
# ===========================================================================

#: Gernots Fall im Kleinen: Ware zu 7 %, Pfandkiste zu 19 %.
#: 2,50 × 7 % = 0,175 → 0,18 | 4,50 × 19 % = 0,855 → 0,86 | zusammen 1,04.
#: Einmal über alle Sätze gerundet ergab das 1,03 — 1 ct neben dem Steuerblock.
S2_ZEILEN = [
    {"description": "Erbsen-Schale", "quantity": 1, "unit": "STK",
     "unit_price": 2.50, "tax_rate": "REDUZIERT"},
    {"description": "Pfandkiste IFCO", "quantity": 1, "unit": "STK",
     "unit_price": 4.50, "tax_rate": "STANDARD"},
]


def _s2_kunde(client):
    r = client.post("/api/v1/sales/customers", json={"name": "Oekoring Test", "typ": "HANDEL"})
    assert r.status_code == 201, r.text
    return r.json()


def _s2_rechnung_zeilenweise(client, zeilen=S2_ZEILEN, **kopf):
    """Anlage wie in der Oberfläche: Kopf, dann jede Position einzeln."""
    r = client.post("/api/v1/invoices", json={
        "customer_id": _s2_kunde(client)["id"],
        "invoice_date": date.today().isoformat(),
        **kopf,
    })
    assert r.status_code == 201, r.text
    rechnung = r.json()
    for zeile in zeilen:
        z = client.post(f"/api/v1/invoices/{rechnung['id']}/lines", json=zeile)
        assert z.status_code == 201, z.text
    return client.get(f"/api/v1/invoices/{rechnung['id']}").json()


def _s2_steuerblock(invoice_id):
    """(get_tax_summary(), subtotal, tax_amount, total) direkt aus der DB."""
    from app.models.invoice import Invoice
    with TestingSessionLocal() as db:
        inv = db.get(Invoice, uuid.UUID(invoice_id))
        return inv.get_tax_summary(), inv.subtotal, inv.tax_amount, inv.total


class TestS2Rechenregel:
    """Die eine Rechenregel, ohne Datenbank geprüft."""

    @staticmethod
    def _zeile(betrag, satz):
        from app.models.enums import TaxRate
        return SimpleNamespace(line_total=Decimal(betrag), tax_rate=TaxRate(satz))

    def test_steuer_wird_je_satz_gerundet(self):
        from app.models.invoice import steuer_je_satz
        saetze = steuer_je_satz(
            [self._zeile("4.50", "STANDARD"), self._zeile("2.50", "REDUZIERT")], Decimal("0"))

        # aufsteigend nach Satz, unabhängig von der Positionsreihenfolge
        assert [(s["percent"], s["base"], s["tax"]) for s in saetze] == [
            (7, Decimal("2.50"), Decimal("0.18")),
            (19, Decimal("4.50"), Decimal("0.86")),
        ]

    def test_rabatt_wird_je_satz_gerundet(self):
        """3,8 % auf 25,00 € (7 %) und 9,00 € (19 %)."""
        from app.models.invoice import steuer_je_satz
        saetze = steuer_je_satz(
            [self._zeile("25.00", "REDUZIERT"), self._zeile("9.00", "STANDARD")], Decimal("3.8"))

        assert [(s["rabatt"], s["base"], s["tax"]) for s in saetze] == [
            (Decimal("0.95"), Decimal("24.05"), Decimal("1.68")),
            (Decimal("0.34"), Decimal("8.66"), Decimal("1.65")),
        ]


class TestS2Rechnungssummen:

    def test_summe_der_satzsteuern_ist_tax_amount(self, client):
        """Bisher: tax_amount 1,03, Steuerblock 0,18 + 0,86 = 1,04."""
        rechnung = _s2_rechnung_zeilenweise(client)

        saetze, netto, steuer, brutto = _s2_steuerblock(rechnung["id"])
        assert sum(s["tax"] for s in saetze) == steuer == Decimal("1.04")
        assert sum(s["base"] for s in saetze) == netto == Decimal("7.00")
        assert brutto == Decimal("8.04")
        assert Decimal(str(rechnung["tax_amount"])) == Decimal("1.04")

    def test_mit_rechnungsrabatt_gehen_alle_summen_auf(self, client):
        rechnung = _s2_rechnung_zeilenweise(client, zeilen=[
            {"description": "Erbsen-Schale", "quantity": 10, "unit": "STK",
             "unit_price": 2.50, "tax_rate": "REDUZIERT"},
            {"description": "Pfandkiste IFCO", "quantity": 2, "unit": "STK",
             "unit_price": 4.50, "tax_rate": "STANDARD"},
        ], discount_percent=3.8)

        saetze, netto, steuer, brutto = _s2_steuerblock(rechnung["id"])
        assert Decimal(str(rechnung["discount_amount"])) == sum(s["rabatt"] for s in saetze) == Decimal("1.29")
        assert netto == sum(s["base"] for s in saetze) == Decimal("32.71")
        assert steuer == sum(s["tax"] for s in saetze) == Decimal("3.33")
        assert brutto == netto + steuer == Decimal("36.04")

    def test_rabatt_entfernen_setzt_rabattbetrag_zurueck(self, client):
        """Sonst druckt das PDF 'Zwischensumme' und 'Rabatt' mit dem alten Betrag."""
        rechnung = _s2_rechnung_zeilenweise(client, discount_percent=10)
        assert Decimal(str(rechnung["discount_amount"])) == Decimal("0.70")

        r = client.patch(f"/api/v1/invoices/{rechnung['id']}", json={"discount_percent": 0})

        assert r.status_code == 200, r.text
        assert Decimal(str(r.json()["discount_amount"])) == Decimal("0")
        assert Decimal(str(r.json()["total"])) == Decimal("8.04")


class TestS2KeinAbrufSchreibt:

    def test_get_tax_summary_liest_gespeicherte_zeilenbetraege(self, client):
        """PDF-Abruf und DATEV lesen get_tax_summary() auch für versendete
        Rechnungen. Es darf line_total nicht aus Menge × Preis neu rechnen
        und nichts an der Session verändern."""
        from app.models.invoice import Invoice
        rechnung = _s2_rechnung_zeilenweise(client)
        # Gespeicherten Zeilenbetrag abweichend von Menge × Preis setzen, damit
        # ein Neuberechnen auffällt. Nur im Test — echte Rechnungen nie direkt ändern.
        with TestingSessionLocal() as db:
            inv = db.get(Invoice, uuid.UUID(rechnung["id"]))
            next(l for l in inv.lines if l.description == "Erbsen-Schale").line_total = Decimal("2.40")
            db.commit()

        with TestingSessionLocal() as db:
            inv = db.get(Invoice, uuid.UUID(rechnung["id"]))
            saetze = inv.get_tax_summary()

            assert not db.dirty, f"get_tax_summary() hat verändert: {db.dirty}"
            assert [(s["percent"], s["base"]) for s in saetze] == [
                (7, Decimal("2.40")),
                (19, Decimal("4.50")),
            ]


# --- S2: POST /invoices mit Positionen, recalculate_totals (Task 9) --------

def _s2_rechnung_in_einem_aufruf(client, zeilen=S2_ZEILEN, **kopf):
    """POST /invoices mit Positionen — der Weg, der bisher nur Zeile 1 zählte."""
    r = client.post("/api/v1/invoices", json={
        "customer_id": _s2_kunde(client)["id"],
        "invoice_date": date.today().isoformat(),
        "lines": zeilen,
        **kopf,
    })
    assert r.status_code == 201, r.text
    return r.json()


class TestS2AnlageMitPositionen:

    def test_anlage_mit_positionen_zaehlt_alle_zeilen(self, client):
        """Bisher 2,50 / 0,18 / 2,68 — nur die erste Position."""
        rechnung = _s2_rechnung_in_einem_aufruf(client)

        assert Decimal(str(rechnung["subtotal"])) == Decimal("7.00")
        assert Decimal(str(rechnung["tax_amount"])) == Decimal("1.04")
        assert Decimal(str(rechnung["total"])) == Decimal("8.04")
        detail = client.get(f"/api/v1/invoices/{rechnung['id']}").json()
        assert len(detail["lines"]) == 2

    def test_hilfsfunktion_rechnet_ohne_geloeschte_position(self, client):
        """recalculate_totals sieht auch eine gelöschte Position nicht mehr —
        Task 25 nutzt das beim Löschen einer Entwurfsposition."""
        from app.models.invoice import Invoice
        from app.services.invoice_service import InvoiceService

        rechnung = _s2_rechnung_in_einem_aufruf(client)
        with TestingSessionLocal() as db:
            inv = db.get(Invoice, uuid.UUID(rechnung["id"]))
            assert len(inv.lines) == 2  # Collection ist jetzt geladen
            db.delete(inv.lines[1])
            InvoiceService(db).recalculate_totals(inv)
            assert len(inv.lines) == 1
            # übrig: 2,50 € zu 7 % = 2,50 netto + 0,18 USt
            assert inv.total == Decimal("2.68")


# --- S2 Task 3 -------------------------------------------------------------

def _s2_pdf_texte(client, invoice_id) -> list[str]:
    """Alle Textstücke des Rechnungs-PDFs, je Tabellenzelle eins.

    ReportLab schreibt jede Zelle als '(<text>) Tj'; '€' erscheint dabei als
    Oktal-Escape '\\200'. Deshalb prüfen die Tests mit startswith.
    """
    r = client.get(f"/api/v1/invoices/{invoice_id}/pdf")
    assert r.status_code == 200, r.text
    roh = _pdf_text(r.content).decode("latin-1", errors="ignore")
    return re.findall(r"\((.*?)\) Tj", roh)


class TestS2RechnungsPdf:

    def test_pdf_weist_entgelt_und_steuer_je_satz_aus(self, client):
        rechnung = _s2_rechnung_in_einem_aufruf(client)

        texte = _s2_pdf_texte(client, rechnung["id"])

        assert "MwSt" in texte, "Spalte MwSt fehlt in der Positionstabelle"
        assert "7 %" in texte and "19 %" in texte, "Satz je Position fehlt"
        assert any(t.startswith("USt 7 % auf 2.50 ") for t in texte), texte
        assert any(t.startswith("USt 19 % auf 4.50 ") for t in texte), texte
        assert any(t.startswith("0.18 ") for t in texte)
        assert any(t.startswith("0.86 ") for t in texte)
        assert any(t.startswith("8.04 ") for t in texte)
        assert "USt:" not in texte, "Pauschale USt-Zeile statt Ausweis je Satz"

    def test_ein_satz_ergibt_eine_steuerzeile(self, client):
        rechnung = _s2_rechnung_in_einem_aufruf(client, zeilen=[S2_ZEILEN[0]])

        texte = _s2_pdf_texte(client, rechnung["id"])

        assert any(t.startswith("USt 7 % auf 2.50 ") for t in texte), texte
        assert not any(t.startswith("USt 19 %") for t in texte)

    def test_altrechnung_behaelt_festgeschriebene_betraege(self, client):
        """GoBD: eine vor der Umstellung versendete Rechnung mit 1,03 € USt
        darf beim erneuten Abruf nicht plötzlich 1,04 € zeigen."""
        from app.models.invoice import Invoice
        rechnung = _s2_rechnung_in_einem_aufruf(client)
        assert client.post(f"/api/v1/invoices/{rechnung['id']}/finalize").status_code == 200
        # Zustand einer Altrechnung nachstellen: Summen mit der alten Rundung.
        # Nur im Test — in echten Daten werden Rechnungen nie direkt geändert.
        with TestingSessionLocal() as db:
            inv = db.get(Invoice, uuid.UUID(rechnung["id"]))
            inv.tax_amount = Decimal("1.03")
            inv.total = Decimal("8.03")
            db.commit()

        texte = _s2_pdf_texte(client, rechnung["id"])

        assert "USt:" in texte
        assert any(t.startswith("1.03 ") for t in texte)
        assert any(t.startswith("8.03 ") for t in texte)
        assert not any(t.startswith("USt 7 % auf") for t in texte)
        # Die DB behält die festgeschriebenen Summen. GET /pdf committet nie;
        # dass get_tax_summary() selbst nichts verändert, prüft
        # TestS2KeinAbrufSchreibt (Task 8).
        _, _, steuer, brutto = _s2_steuerblock(rechnung["id"])
        assert (steuer, brutto) == (Decimal("1.03"), Decimal("8.03"))

    def test_sammelrechnung_pdf_laesst_sich_erzeugen(self, client, sample_customer):
        """Die Lieferschein-Anhangstabelle nutzte Decimal ohne Import → 500."""
        from tests.test_sammelrechnung import COMMIT, _bestellung_mit_ls
        _bestellung_mit_ls(client, sample_customer["id"], "2026-03-05", [
            {"product_name": "Erbsen-Schale", "quantity": 10, "unit": "STK",
             "unit_price": 2.50, "tax_rate": "REDUZIERT"},
        ])
        r = client.post(COMMIT, json={"period_from": "2026-03-01", "period_to": "2026-03-31"})
        assert r.status_code == 201, r.text

        texte = _s2_pdf_texte(client, r.json()["rechnungen"][0]["id"])

        assert "Enthaltene Lieferscheine" in texte
        assert any(t.startswith("USt 7 % auf 25.00 ") for t in texte), texte


# --- S2 Task 4 -------------------------------------------------------------

def _s2_texte(pdf_bytes) -> list[str]:
    """Wie _s2_pdf_texte, aber für PDF-Bytes, die nicht über GET /pdf kommen."""
    return re.findall(r"\((.*?)\) Tj", _pdf_text(pdf_bytes).decode("latin-1", errors="ignore"))


def _s2_alte_summen_setzen(invoice_id):
    """Summen mit der früheren Rundung (1,03 / 8,03) nachstellen.
    Nur im Test — in echten Daten werden Rechnungen nie direkt geändert."""
    from app.models.invoice import Invoice
    with TestingSessionLocal() as db:
        inv = db.get(Invoice, uuid.UUID(invoice_id))
        inv.tax_amount = Decimal("1.03")
        inv.total = Decimal("8.03")
        db.commit()


def _s2_mailen(client, monkeypatch, invoice_id) -> dict:
    """POST /invoices/{id}/send mit abgefangenem Versand; liefert die Mail-Argumente."""
    versendet = {}
    monkeypatch.setattr("app.api.v1.invoices.send_email", lambda **kw: versendet.update(kw))
    r = client.post(f"/api/v1/invoices/{invoice_id}/send",
                    params={"to_email": "einkauf@oekoring.example"})
    assert r.status_code == 200, r.text
    return versendet


class TestS2Mailversand:

    def test_mailen_eines_altentwurfs_rechnet_vorher_neu(self, client, monkeypatch):
        """'Mailen' stellt einen Entwurf aus (ENTWURF -> OFFEN), genau wie
        /finalize. Ein Entwurf mit Summen aus der alten Rundung ging bisher mit
        1,03 € und der pauschalen Zeile 'USt:' hinaus."""
        rechnung = _s2_rechnung_in_einem_aufruf(client)
        _s2_alte_summen_setzen(rechnung["id"])

        mail = _s2_mailen(client, monkeypatch, rechnung["id"])

        detail = client.get(f"/api/v1/invoices/{rechnung['id']}").json()
        assert detail["status"] == "OFFEN"
        assert Decimal(str(detail["tax_amount"])) == Decimal("1.04")
        assert Decimal(str(detail["total"])) == Decimal("8.04")
        texte = _s2_texte(mail["attachment_bytes"])
        assert any(t.startswith("USt 7 % auf 2.50 ") for t in texte), texte
        assert any(t.startswith("USt 19 % auf 4.50 ") for t in texte), texte
        assert "USt:" not in texte
        assert "8.04 EUR" in mail["body"]

    def test_mailen_einer_versendeten_rechnung_rechnet_nicht_neu(self, client, monkeypatch):
        """GoBD: erneutes Mailen einer festgeschriebenen Rechnung ändert keine Beträge."""
        rechnung = _s2_rechnung_in_einem_aufruf(client)
        assert client.post(f"/api/v1/invoices/{rechnung['id']}/finalize").status_code == 200
        _s2_alte_summen_setzen(rechnung["id"])

        mail = _s2_mailen(client, monkeypatch, rechnung["id"])

        _, _, steuer, brutto = _s2_steuerblock(rechnung["id"])
        assert (steuer, brutto) == (Decimal("1.03"), Decimal("8.03"))
        texte = _s2_texte(mail["attachment_bytes"])
        assert "USt:" in texte
        assert any(t.startswith("1.03 ") for t in texte)
        assert "8.03 EUR" in mail["body"]


# ---------------------------------------------------------------------------
# S3: Storno mehrzeiliger Rechnungen, Ausgleich, Warnungen
# ---------------------------------------------------------------------------

_S3_ZEILEN = [
    {"description": "BIO Erbsen-Schale", "quantity": 10, "unit": "STK",
     "unit_price": 1.00, "tax_rate": "REDUZIERT"},
    {"description": "Versandkarton", "quantity": 1, "unit": "STK",
     "unit_price": 6.67, "tax_rate": "STANDARD"},
]


def _s3_pfandartikel(client):
    """Pfandartikel wie im Produktstamm von Minga: 19 %, als Pfand gekennzeichnet."""
    from app.models.unit import UnitOfMeasure, UnitCategory
    with TestingSessionLocal() as db:
        unit = db.query(UnitOfMeasure).filter_by(code="STK").first()
        if unit is None:
            unit = UnitOfMeasure(code="STK", name="Stück", category=UnitCategory.COUNT)
            db.add(unit)
            db.commit()
        unit_id = str(unit.id)
    r = client.post("/api/v1/products", json={
        "sku": f"IFCO-{uuid.uuid4().hex[:6]}", "name": "Pfand IFCO-Kiste",
        "category": "PFAND", "base_price": "6.67", "base_unit_id": unit_id,
        "tax_rate": "STANDARD", "is_deposit": True,
    })
    assert r.status_code in (200, 201), r.text
    return r.json()


def _s3_altrechnung(client, kunde, *, rabatt=None, lieferdatum=None, konto_ware=None):
    """Stellt RE-2026-00002/4 nach: Ware zu 7 % plus Pfand-Produktzeile, die
    auf der Rechnung mit 7 % statt 19 % steht — versendet, also OFFEN.

    Die Pfandzeile wird direkt angelegt: so steht sie in Produktion, und die
    API muss einen vom Produktstamm abweichenden Satz nicht mehr annehmen.
    konto_ware: ausdrücklich gesetztes Sonderkonto der Warenzeile
    (InvoiceLineCreate.buchungskonto), sonst das Standardkonto zum Satz.
    Ohne Rabatt: 16,67 netto, 1,17 USt, 17,84 brutto.
    """
    from app.models.invoice import InvoiceLine, TaxRate

    pfand = _s3_pfandartikel(client)
    body = {"customer_id": kunde["id"], "invoice_date": date.today().isoformat()}
    if lieferdatum:
        body["delivery_date"] = lieferdatum
    r = client.post("/api/v1/invoices", json=body)
    assert r.status_code == 201, r.text
    rechnung = r.json()

    ware = dict(_S3_ZEILEN[0])
    if konto_ware:
        ware["buchungskonto"] = konto_ware
    r = client.post(f"/api/v1/invoices/{rechnung['id']}/lines", json=ware)
    assert r.status_code == 201, r.text
    with TestingSessionLocal() as db:
        db.add(InvoiceLine(
            invoice_id=uuid.UUID(rechnung["id"]), position=2,
            product_id=uuid.UUID(pfand["id"]), sku=pfand["sku"],
            description="Pfand IFCO-Kiste",
            quantity=Decimal("1"), unit="STK", unit_price=Decimal("6.67"),
            tax_rate=TaxRate.REDUZIERT, is_deposit=True, buchungskonto="8300",
        ))
        db.commit()

    if rabatt is not None:
        r = client.patch(f"/api/v1/invoices/{rechnung['id']}", json={"discount_percent": rabatt})
        assert r.status_code == 200, r.text
    # finalize rechnet die Summen in einem frischen Request über alle Zeilen
    r = client.post(f"/api/v1/invoices/{rechnung['id']}/finalize")
    assert r.status_code == 200, r.text
    return r.json()


def _s3_storniere(client, rechnung, **extra):
    body = {"reason": "Pfand mit 7 % statt 19 %", "reason_code": "SONSTIGES", **extra}
    return client.post(f"/api/v1/invoices/{rechnung['id']}/cancel", json=body)


def _s3_betrag(wert):
    # Beträge kommen als JSON-String ("17.84") — nie über float vergleichen
    return Decimal(str(wert))


class TestS3StornoMehrzeilig:
    """Der Storno summierte bei mehreren Positionen nur die erste (−10,70 statt −17,84)."""

    def test_storno_summiert_alle_positionen(self, client, sample_customer):
        original = _s3_altrechnung(client, sample_customer)
        assert _s3_betrag(original["total"]) == Decimal("17.84")

        r = _s3_storniere(client, original)
        assert r.status_code == 200, r.text
        storno = r.json()["credit_note"]

        assert _s3_betrag(storno["subtotal"]) == Decimal("-16.67")
        assert _s3_betrag(storno["tax_amount"]) == Decimal("-1.17")
        assert _s3_betrag(storno["total"]) == Decimal("-17.84")

    def test_storno_ist_spiegelbild_je_position(self, client, sample_customer):
        """Satz, Konto, Pfandkennzeichen und Artikelnummer kommen von der
        Originalzeile. Die Warenzeile steht auf einem Sonderkonto (8338): der
        frühere Weg über add_line setzte dort das Standardkonto zum Satz (8300),
        die Gegenbuchung landete auf einem anderen Konto als die Buchung."""
        from app.models.invoice import Invoice

        original = _s3_altrechnung(client, sample_customer, konto_ware="8338")
        storno = _s3_storniere(client, original).json()["credit_note"]

        orig = client.get(f"/api/v1/invoices/{original['id']}").json()
        sto = client.get(f"/api/v1/invoices/{storno['id']}").json()
        assert len(sto["lines"]) == len(orig["lines"]) == 2
        for o, s in zip(orig["lines"], sto["lines"]):
            assert s["position"] == o["position"]
            assert s["tax_rate"] == o["tax_rate"]
            assert s["buchungskonto"] == o["buchungskonto"]
            assert s["product_id"] == o["product_id"]
            assert s["sku"] == o["sku"]
            assert _s3_betrag(s["discount_percent"]) == _s3_betrag(o["discount_percent"])
            assert _s3_betrag(s["quantity"]) == -_s3_betrag(o["quantity"])
            assert _s3_betrag(s["line_total"]) == -_s3_betrag(o["line_total"])
        assert _s3_betrag(sto["total_deposit"]) == -_s3_betrag(orig["total_deposit"])

        # is_deposit steht erst ab S5 in der Zeilenantwort — deshalb über das ORM
        with TestingSessionLocal() as db:
            o_zeilen = db.get(Invoice, uuid.UUID(original["id"])).lines
            s_zeilen = db.get(Invoice, uuid.UUID(storno["id"])).lines
            assert [z.is_deposit for z in s_zeilen] == [z.is_deposit for z in o_zeilen] == [False, True]

        # GoBD: das Original bleibt unverändert
        assert _s3_betrag(orig["total"]) == Decimal("17.84")
        assert [l["tax_rate"] for l in orig["lines"]] == ["REDUZIERT", "REDUZIERT"]
        assert [l["buchungskonto"] for l in orig["lines"]] == ["8338", "8300"]

    def test_storno_mit_zwei_saetzen_spiegelt_steuer_je_satz(self, client, sample_customer):
        """Stornorechnung mit 7 % und 19 %: je Satz heben Entgelt und Steuer
        das Original genau auf, und die Satzsteuern ergeben tax_amount
        (dieselbe Prüfung wie steuerausweis_stimmt aus S2)."""
        from app.models.invoice import Invoice

        r = client.post("/api/v1/invoices", json={
            "customer_id": sample_customer["id"],
            "invoice_date": date.today().isoformat(),
        })
        assert r.status_code == 201, r.text
        rechnung = r.json()
        # Positionen einzeln: jeder Request hat eine frische Session, die
        # Summen des Originals stimmen also auch ohne den Fix in add_line.
        for zeile in _S3_ZEILEN:
            r = client.post(f"/api/v1/invoices/{rechnung['id']}/lines", json=zeile)
            assert r.status_code == 201, r.text
        r = client.post(f"/api/v1/invoices/{rechnung['id']}/finalize")
        assert r.status_code == 200, r.text
        original = r.json()
        # 10,00 × 7 % + 6,67 × 19 % = 16,67 netto + 0,70 + 1,27 USt
        assert _s3_betrag(original["total"]) == Decimal("18.64")

        storno = _s3_storniere(client, original).json()["credit_note"]

        assert _s3_betrag(storno["subtotal"]) == Decimal("-16.67")
        assert _s3_betrag(storno["tax_amount"]) == Decimal("-1.97")
        assert _s3_betrag(storno["total"]) == Decimal("-18.64")
        with TestingSessionLocal() as db:
            o = db.get(Invoice, uuid.UUID(original["id"]))
            s = db.get(Invoice, uuid.UUID(storno["id"]))
            je_satz_o = {e["rate"]: (e["base"], e["tax"]) for e in o.get_tax_summary()}
            je_satz_s = {e["rate"]: (e["base"], e["tax"]) for e in s.get_tax_summary()}
            assert sum(b for b, _ in je_satz_s.values()) == s.subtotal
            assert sum(t for _, t in je_satz_s.values()) == s.tax_amount
        assert len(je_satz_s) == 2
        assert je_satz_s == {satz: (-b, -t) for satz, (b, t) in je_satz_o.items()}

    def test_storno_uebernimmt_rabatt_und_lieferdatum(self, client, sample_customer):
        original = _s3_altrechnung(client, sample_customer, rabatt=10, lieferdatum="2026-10-07")
        storno = _s3_storniere(client, original).json()["credit_note"]

        assert _s3_betrag(storno["discount_percent"]) == Decimal("10")
        assert _s3_betrag(storno["total"]) == -_s3_betrag(original["total"])
        assert storno["delivery_date"] == "2026-10-07"


class TestS3StornoAusgeglichen:
    """Stornorechnung und Original gleichen sich aus und fallen aus
    Überfälligkeit, Mahnlauf, offenen Posten und Umsatz heraus."""

    def test_beide_belege_stehen_auf_storniert(self, client, sample_customer):
        original = _s3_altrechnung(client, sample_customer)
        antwort = _s3_storniere(client, original).json()

        assert antwort["invoice"]["status"] == "STORNIERT"
        storno = antwort["credit_note"]
        assert storno["status"] == "STORNIERT"
        assert storno["due_date"] == storno["invoice_date"]

    def _faelligkeit_in_der_vergangenheit(self, *ids):
        from app.models.invoice import Invoice
        with TestingSessionLocal() as db:
            for i in ids:
                db.get(Invoice, uuid.UUID(i)).due_date = date.today() - timedelta(days=30)
            db.commit()

    def _als_altbestand_offen(self, storno_id):
        """Stornorechnung so, wie der alte Code sie hinterließ: OFFEN, fällig."""
        from app.models.invoice import Invoice, InvoiceStatus
        with TestingSessionLocal() as db:
            alt = db.get(Invoice, uuid.UUID(storno_id))
            alt.status = InvoiceStatus.OFFEN
            alt.due_date = date.today() - timedelta(days=30)
            db.commit()

    def test_storno_ist_nie_ueberfaellig(self, client, sample_customer):
        original = _s3_altrechnung(client, sample_customer)
        storno = _s3_storniere(client, original).json()["credit_note"]
        self._faelligkeit_in_der_vergangenheit(original["id"], storno["id"])

        r = client.get("/api/v1/invoices/overdue")
        assert r.status_code == 200, r.text
        ids = {i["id"] for i in r.json()}
        assert original["id"] not in ids
        assert storno["id"] not in ids
        assert client.get(f"/api/v1/invoices/{storno['id']}").json()["status"] == "STORNIERT"

    def test_ueberfaellig_liste_ignoriert_alten_storno_auf_offen(self, client, sample_customer):
        """Altbestand über die API: GET /invoices/overdue ruft
        InvoiceService.check_overdue_invoices und schreibt UEBERFAELLIG."""
        original = _s3_altrechnung(client, sample_customer)
        storno = _s3_storniere(client, original).json()["credit_note"]
        self._als_altbestand_offen(storno["id"])

        r = client.get("/api/v1/invoices/overdue")
        assert r.status_code == 200, r.text
        assert storno["id"] not in {i["id"] for i in r.json()}
        assert client.get(f"/api/v1/invoices/{storno['id']}").json()["status"] == "OFFEN"

    def test_geplante_jobs_ignorieren_negative_belege(self, client, sample_customer):
        """Altbestand: eine Stornorechnung, die der alte Code auf OFFEN gesetzt hat,
        wird weder überfällig noch gemahnt."""
        from app.models.invoice import Invoice, InvoiceStatus
        from app.tasks.invoice_tasks import check_overdue_invoices, send_payment_reminders

        original = _s3_altrechnung(client, sample_customer)
        storno = _s3_storniere(client, original).json()["credit_note"]
        self._als_altbestand_offen(storno["id"])

        with patch("app.tasks.invoice_tasks.SessionLocal", return_value=TestingSessionLocal()):
            assert check_overdue_invoices()["newly_overdue"] == 0
        assert client.get(f"/api/v1/invoices/{storno['id']}").json()["status"] == "OFFEN"

        with TestingSessionLocal() as db:
            db.get(Invoice, uuid.UUID(storno["id"])).status = InvoiceStatus.UEBERFAELLIG
            db.commit()
        mail = MagicMock()
        mail.send_email.return_value = True
        with patch("app.tasks.invoice_tasks.email_service", mail), \
             patch("app.tasks.invoice_tasks.SessionLocal", return_value=TestingSessionLocal()):
            assert send_payment_reminders()["reminders_sent"] == 0
        mail.send_email.assert_not_called()

    def test_weder_offener_posten_noch_umsatz(self, client, sample_customer):
        original = _s3_altrechnung(client, sample_customer)
        storno = _s3_storniere(client, original).json()["credit_note"]

        offen = {i["id"] for i in client.get("/api/v1/invoices", params={"status": "OFFEN"}).json()}
        assert storno["id"] not in offen and original["id"] not in offen

        heute = date.today().isoformat()
        summe = client.get("/api/v1/invoices/revenue-summary",
                           params={"from_date": heute, "to_date": heute}).json()
        # Paar hebt sich auf: kein Umsatz, nichts offen — vorher −17,84 Umsatz
        assert _s3_betrag(summe["total_revenue"]) == Decimal("0")
        assert _s3_betrag(summe["open_amount"]) == Decimal("0")

    def test_stornorechnung_kann_versendet_werden(self, client, sample_customer):
        original = _s3_altrechnung(client, sample_customer)
        storno = _s3_storniere(client, original).json()["credit_note"]

        gesendet = {}
        def _fake_send_email(**kwargs):
            gesendet.update(kwargs)
        with patch("app.api.v1.invoices.send_email", _fake_send_email):
            r = client.post(f"/api/v1/invoices/{storno['id']}/send",
                            params={"to_email": "kunde@example.com"})
            assert r.status_code == 200, r.text
            assert "Stornorechnung" in gesendet["subject"]
            assert original["invoice_number"] in gesendet["body"]
            assert "Fällig" not in gesendet["body"]
            assert gesendet["attachment_bytes"].startswith(b"%PDF")

            # das stornierte Original bleibt gesperrt
            r = client.post(f"/api/v1/invoices/{original['id']}/send",
                            params={"to_email": "kunde@example.com"})
            assert r.status_code == 400, r.text

        assert client.get(f"/api/v1/invoices/{storno['id']}").json()["status"] == "STORNIERT"

    def test_entwurf_bekommt_keine_stornorechnung(self, client, sample_customer):
        r = client.post("/api/v1/invoices", json={
            "customer_id": sample_customer["id"],
            "invoice_date": date.today().isoformat(),
            "lines": _S3_ZEILEN,
        })
        entwurf = r.json()

        r = _s3_storniere(client, entwurf)
        assert r.status_code == 400, r.text
        assert "Entwurf" in r.json()["detail"]

        r = _s3_storniere(client, entwurf, create_credit_note=False)
        assert r.status_code == 200, r.text
        assert r.json()["credit_note"] is None


class TestS3StornoWarnungenUndGrund:
    def test_storno_einer_teilbezahlten_rechnung_warnt(self, client, sample_customer):
        original = _s3_altrechnung(client, sample_customer)
        r = client.post(f"/api/v1/invoices/{original['id']}/payments", json={
            "invoice_id": original["id"], "payment_date": date.today().isoformat(),
            "amount": "5.00",
        })
        assert r.status_code == 201, r.text

        antwort = _s3_storniere(client, original).json()

        assert len(antwort["warnungen"]) == 1
        assert "5,00 €" in antwort["warnungen"][0]
        assert original["invoice_number"] in antwort["warnungen"][0]
        detail = client.get(f"/api/v1/invoices/{original['id']}").json()
        assert "bereits gezahlt: 5,00 €" in detail["internal_notes"]
        assert _s3_betrag(detail["paid_amount"]) == Decimal("5.00")

    def test_storno_ohne_zahlung_ohne_warnung(self, client, sample_customer):
        original = _s3_altrechnung(client, sample_customer)
        assert _s3_storniere(client, original).json()["warnungen"] == []

    def test_kopie_in_lexoffice_wird_gemeldet(self, client, sample_customer):
        from app.models.invoice import Invoice
        original = _s3_altrechnung(client, sample_customer)
        with TestingSessionLocal() as db:
            db.get(Invoice, uuid.UUID(original["id"])).lexoffice_id = "lex-123"
            db.commit()

        warnungen = _s3_storniere(client, original).json()["warnungen"]

        assert any("lexoffice" in w for w in warnungen)

    def test_stornogrund_falscher_steuersatz(self, client, sample_customer):
        original = _s3_altrechnung(client, sample_customer)

        r = _s3_storniere(client, original, reason_code="FALSCHER_STEUERSATZ",
                          reason="Pfand mit 7 % statt 19 % berechnet")
        assert r.status_code == 200, r.text

        orig = client.get(f"/api/v1/invoices/{original['id']}").json()
        sto = client.get(f"/api/v1/invoices/{r.json()['credit_note']['id']}").json()
        assert "[FALSCHER_STEUERSATZ]" in orig["internal_notes"]
        assert "[FALSCHER_STEUERSATZ]" in sto["internal_notes"]


# ---------------------------------------------------------------------------
# S4: DATEV-Export
# ---------------------------------------------------------------------------

DATEV_KOPF = [
    "Umsatz", "Soll/Haben", "WKZ", "Kurs", "Basisumsatz",
    "Konto", "Gegenkonto", "BU-Schlüssel", "Belegdatum",
    "Belegfeld 1", "Belegfeld 2", "Buchungstext",
]

#: Gemischte Rechnung wie RE-00002: Ware zu 7 %, Pfandkiste zu 19 %.
#: 10 × 2,50 = 25,00 netto + 1,75 USt = 26,75 | 2 × 3,00 = 6,00 netto + 1,14 USt = 7,14
DATEV_GEMISCHT = [
    ("Erbsen-Schale", 10, "2.50", "REDUZIERT"),
    ("Pfandkiste 6er", 2, "3.00", "STANDARD"),
]


def _datev_kunde(client, name="Ökoring Testkunde", konto="10008", rabatt=None):
    body = {"name": name, "typ": "HANDEL"}
    if konto is not None:
        body["datev_account"] = konto
    if rabatt is not None:
        body["discount_percent"] = rabatt
    r = client.post("/api/v1/sales/customers", json=body)
    assert r.status_code == 201, r.text
    return r.json()


def _datev_rechnung(client, kunde, positionen, typ="RECHNUNG", finalisieren=True, **kopf):
    """positionen: (Beschreibung, Menge, Preis, Satz[, Sonderkonto])."""
    r = client.post("/api/v1/invoices", json={
        "customer_id": kunde["id"], "invoice_date": date.today().isoformat(),
        "invoice_type": typ, **kopf,
    })
    assert r.status_code == 201, r.text
    rechnung = r.json()
    for beschreibung, menge, preis, satz, *sonderkonto in positionen:
        zeile = {
            "description": beschreibung, "quantity": menge, "unit": "STK",
            "unit_price": preis, "tax_rate": satz,
        }
        if sonderkonto:
            zeile["buchungskonto"] = sonderkonto[0]
        r = client.post(f"/api/v1/invoices/{rechnung['id']}/lines", json=zeile)
        assert r.status_code == 201, r.text
    if finalisieren:
        r = client.post(f"/api/v1/invoices/{rechnung['id']}/finalize")
        assert r.status_code == 200, r.text
        rechnung = r.json()
    return rechnung


def _datev_export(client, erneut=None, zahlungen=False):
    body = {
        "from_date": date.today().isoformat(), "to_date": date.today().isoformat(),
        "include_payments": zahlungen,
    }
    if erneut is not None:
        body["erneut_exportieren"] = erneut
    r = client.post("/api/v1/invoices/datev-export", json=body)
    assert r.status_code == 200, r.text
    daten = r.json()
    zeilen = list(csv.reader(io.StringIO(daten["csv_content"]), delimiter=";"))
    return daten, zeilen[0], zeilen[1:]


def _datev_zeile(betrag, sh, konto, gegenkonto, belegnr, text):
    """Eine erwartete Buchungszeile, Spalte für Spalte in Kopf-Reihenfolge."""
    return [betrag, sh, "EUR", "", "", konto, gegenkonto, "",
            date.today().strftime("%d%m"), belegnr, "", text]


class TestDatevKontierung:
    """Je Rechnung und Erlöskonto eine Zeile: Debitor an Erlös, Richtung über S/H."""

    def test_gemischte_rechnung_spaltengenau(self, client):
        kunde = _datev_kunde(client)
        rechnung = _datev_rechnung(client, kunde, DATEV_GEMISCHT)
        nr = rechnung["invoice_number"]

        daten, kopf, zeilen = _datev_export(client)

        assert kopf == DATEV_KOPF
        assert zeilen == [
            _datev_zeile("26,75", "S", "10008", "8300", nr, "Rechnung 7 % Ökoring Testkunde"),
            _datev_zeile("7,14", "S", "10008", "8400", nr, "Rechnung 19 % Ökoring Testkunde"),
        ]
        assert daten["record_count"] == 2
        assert Decimal(str(daten["total_amount"])) == Decimal("33.89")

    def test_jede_zeile_hat_so_viele_felder_wie_der_kopf(self, client):
        kunde = _datev_kunde(client)
        rechnung = _datev_rechnung(client, kunde, DATEV_GEMISCHT)
        client.post(f"/api/v1/invoices/{rechnung['id']}/payments", json={
            "invoice_id": rechnung["id"], "amount": "33.89",
            "payment_date": date.today().isoformat(),
        })

        _, kopf, zeilen = _datev_export(client, zahlungen=True)

        assert len(zeilen) == 3
        assert all(len(z) == len(kopf) for z in zeilen)

    def test_kopfkonto_ueberschreibt_die_konten_je_satz_nicht(self, client):
        """Altbestand: jede Rechnung trägt im Kopf 8300. Der Export darf es nicht verwenden."""
        kunde = _datev_kunde(client)
        rechnung = _datev_rechnung(client, kunde, DATEV_GEMISCHT, buchungskonto="8300")
        assert rechnung["buchungskonto"] == "8300"

        _, _, zeilen = _datev_export(client)

        assert [z[6] for z in zeilen] == ["8300", "8400"]

    def test_neue_rechnung_ohne_kopfkonto(self, client):
        kunde = _datev_kunde(client)
        rechnung = _datev_rechnung(client, kunde, DATEV_GEMISCHT, finalisieren=False)
        assert rechnung["buchungskonto"] is None

    def test_konto_folgt_dem_steuersatz_der_position(self, client):
        """Ein nachträglich geänderter Satz darf nicht auf dem alten Standardkonto landen."""
        kunde = _datev_kunde(client)
        rechnung = _datev_rechnung(client, kunde, [("Pfandkiste", 1, "10.00", "REDUZIERT")],
                                   finalisieren=False)
        zeile = client.get(f"/api/v1/invoices/{rechnung['id']}").json()["lines"][0]
        assert zeile["buchungskonto"] == "8300"
        r = client.patch(f"/api/v1/invoices/{rechnung['id']}/lines/{zeile['id']}",
                         json={"tax_rate": "STANDARD"})
        assert r.status_code == 200, r.text
        assert client.post(f"/api/v1/invoices/{rechnung['id']}/finalize").status_code == 200

        _, _, zeilen = _datev_export(client)

        assert zeilen == [
            _datev_zeile("11,90", "S", "10008", "8400", rechnung["invoice_number"],
                         "Rechnung 19 % Ökoring Testkunde"),
        ]

    def test_rechnungsrabatt_wie_steuerausweis_der_rechnung(self, client):
        """Brutto je Satz = Entgelt + USt je Satz, so wie die Rechnung sie ausweist."""
        from app.models.invoice import Invoice
        kunde = _datev_kunde(client, rabatt="10")
        rechnung = _datev_rechnung(client, kunde, DATEV_GEMISCHT)
        nr = rechnung["invoice_number"]

        _, _, zeilen = _datev_export(client)

        # 25,00 − 2,50 = 22,50 + 1,58 USt = 24,08 | 6,00 − 0,60 = 5,40 + 1,03 USt = 6,43
        assert zeilen == [
            _datev_zeile("24,08", "S", "10008", "8300", nr, "Rechnung 7 % Ökoring Testkunde"),
            _datev_zeile("6,43", "S", "10008", "8400", nr, "Rechnung 19 % Ökoring Testkunde"),
        ]
        with TestingSessionLocal() as db:
            ausweis = db.get(Invoice, uuid.UUID(rechnung["id"])).get_tax_summary()
        brutto_je_satz = sorted(f"{s['base'] + s['tax']:.2f}".replace(".", ",") for s in ausweis)
        assert sorted(z[0] for z in zeilen) == brutto_je_satz

    def test_rundungsgrenze_rabatt_wie_rechnungsbetrag(self, client):
        """Rechenregel aus S2 an einer Rundungsgrenze: 22,15 € zu 7 %, 10 % Rabatt.
        Rabatt 2,215 → 2,22, Entgelt 19,93, USt 1,3951 → 1,40, zusammen 21,33.
        Den Rabatt ungerundet abziehen ergäbe 19,935 → 19,94 + 1,40 = 21,34 —
        beim Debitor bliebe 1 ct offen."""
        kunde = _datev_kunde(client, rabatt="10")
        rechnung = _datev_rechnung(client, kunde, [("Erbsen-Schale", 1, "22.15", "REDUZIERT")])
        assert Decimal(str(rechnung["total"])) == Decimal("21.33")

        daten, _, zeilen = _datev_export(client)

        assert zeilen == [
            _datev_zeile("21,33", "S", "10008", "8300", rechnung["invoice_number"],
                         "Rechnung 7 % Ökoring Testkunde"),
        ]
        assert Decimal(str(daten["total_amount"])) == Decimal(str(rechnung["total"]))

    def test_sonderkonto_neben_standardkonto_rest_cent(self, client):
        """Teilen sich zwei Konten einen Satz, rundet jedes Konto für sich:
        25,00 → 22,50 + 1,58 = 24,08 und 22,15 → 19,93 + 1,40 = 21,33, zusammen 45,41.
        Der Steuerausweis zu 7 % lautet 47,15 − 4,72 = 42,43 + 2,97 = 45,40.
        Den Rest-Cent trägt die betragsgrößte Gruppe (8300): 24,07."""
        kunde = _datev_kunde(client, rabatt="10")
        rechnung = _datev_rechnung(client, kunde, [
            ("Erbsen-Schale", 10, "2.50", "REDUZIERT"),
            ("Kresse Sonderaktion", 1, "22.15", "REDUZIERT", "8301"),
        ])
        nr = rechnung["invoice_number"]
        assert Decimal(str(rechnung["total"])) == Decimal("45.40")

        daten, _, zeilen = _datev_export(client)

        assert zeilen == [
            _datev_zeile("24,07", "S", "10008", "8300", nr, "Rechnung 7 % Ökoring Testkunde"),
            _datev_zeile("21,33", "S", "10008", "8301", nr, "Rechnung 7 % Ökoring Testkunde"),
        ]
        assert Decimal(str(daten["total_amount"])) == Decimal(str(rechnung["total"]))

    def test_summe_der_zeilen_gleich_rechnungsbetrag(self, client):
        """Der Debitor bekommt in DATEV genau den Rechnungsbetrag — sonst bleibt
        nach Zahlung ein Cent offen. Wächter: schon vor Task 16 grün (der alte
        Export summierte invoice.total); nach Task 16 grün, weil Export und
        calculate_totals dieselbe Rechenregel steuer_je_satz (S2) nutzen."""
        kunde = _datev_kunde(client, rabatt="10")
        rechnung = _datev_rechnung(client, kunde, DATEV_GEMISCHT)

        daten, _, _ = _datev_export(client)

        assert Decimal(str(daten["total_amount"])) == Decimal(str(rechnung["total"]))

    def test_ohne_debitorenkonto_sammeldebitor(self, client):
        kunde = _datev_kunde(client, konto=None)
        _datev_rechnung(client, kunde, [("Kresse", 1, "10.00", "REDUZIERT")])

        _, _, zeilen = _datev_export(client)

        assert [z[5] for z in zeilen] == ["10000"]

    def test_export_aendert_keinen_beleg(self, client):
        """GoBD: der Export liest. Nur die DATEV-Kennzeichen dürfen sich ändern."""
        kunde = _datev_kunde(client)
        rechnung = _datev_rechnung(client, kunde, DATEV_GEMISCHT)
        vorher = client.get(f"/api/v1/invoices/{rechnung['id']}").json()

        _datev_export(client)

        nachher = client.get(f"/api/v1/invoices/{rechnung['id']}").json()
        for feld in ("subtotal", "tax_amount", "total", "status", "invoice_number"):
            assert nachher[feld] == vorher[feld], feld
        assert nachher["lines"] == vorher["lines"]


class TestDatevGutschrift:
    """Gutschriften mindern: H mit positivem Umsatz."""

    def test_storno_spaltengenau(self, client):
        """Original und Stornorechnung sind beide Belege: S und H heben sich auf."""
        kunde = _datev_kunde(client)
        rechnung = _datev_rechnung(client, kunde, DATEV_GEMISCHT)
        r = client.post(f"/api/v1/invoices/{rechnung['id']}/cancel",
                        json={"reason": "Pfand mit 7 % berechnet"})
        assert r.status_code == 200, r.text
        storno = r.json()["credit_note"]
        nr, snr = rechnung["invoice_number"], storno["invoice_number"]

        daten, _, zeilen = _datev_export(client)

        assert zeilen == [
            _datev_zeile("26,75", "S", "10008", "8300", nr, "Rechnung 7 % Ökoring Testkunde"),
            _datev_zeile("7,14", "S", "10008", "8400", nr, "Rechnung 19 % Ökoring Testkunde"),
            _datev_zeile("26,75", "H", "10008", "8300", snr, f"Storno {nr} 7 % Ökoring Testkunde"),
            _datev_zeile("7,14", "H", "10008", "8400", snr, f"Storno {nr} 19 % Ökoring Testkunde"),
        ]
        assert Decimal(str(daten["total_amount"])) == Decimal("0.00")

    def test_storno_spaeter_exportiert_nur_die_stornorechnung(self, client):
        kunde = _datev_kunde(client)
        rechnung = _datev_rechnung(client, kunde, DATEV_GEMISCHT)
        _datev_export(client)
        r = client.post(f"/api/v1/invoices/{rechnung['id']}/cancel", json={"reason": "Preisfehler"})
        snr = r.json()["credit_note"]["invoice_number"]

        _, _, zeilen = _datev_export(client)

        assert [(z[1], z[9]) for z in zeilen] == [("H", snr), ("H", snr)]

    def test_manuelle_gutschrift_mindert(self, client):
        """Gutschrift von Hand (positive Beträge) ist eine Minderung: H."""
        kunde = _datev_kunde(client)
        gs = _datev_rechnung(client, kunde, [("Preisnachlass", 1, "10.00", "REDUZIERT")],
                             typ="GUTSCHRIFT")

        _, _, zeilen = _datev_export(client)

        assert zeilen == [
            _datev_zeile("10,70", "H", "10008", "8300", gs["invoice_number"],
                         "Gutschrift 7 % Ökoring Testkunde"),
        ]


class TestDatevBelegauswahl:
    """Nur Buchungsbelege: kein Entwurf, keine Proforma, kein Storno ohne Gegenbeleg."""

    def test_storniert_ohne_stornorechnung_wird_nicht_exportiert(self, client):
        """Ohne Gegenbeleg würde DATEV Umsatz buchen, den es nie gab."""
        from app.models.invoice import Invoice, InvoiceStatus
        kunde = _datev_kunde(client)
        rechnung = _datev_rechnung(client, kunde, DATEV_GEMISCHT)
        # Zustand wie nach einem Storno ohne Gegenbeleg (create_credit_note=False
        # oder PATCH status) — direkt gesetzt, damit der Test nicht davon abhängt,
        # ob die API diesen Weg künftig noch erlaubt.
        with TestingSessionLocal() as db:
            db.get(Invoice, uuid.UUID(rechnung["id"])).status = InvoiceStatus.STORNIERT
            db.commit()

        daten, _, zeilen = _datev_export(client)

        assert zeilen == []
        assert daten["record_count"] == 0

    def test_verworfene_gutschrift_von_hand_wird_nicht_exportiert(self, client):
        """Eine von Hand angelegte Gutschrift, deren Entwurf per PATCH status=STORNIERT
        verworfen wurde, ist kein Beleg. Nur die Stornorechnung (GUTSCHRIFT mit
        original_invoice_id) wird unabhängig vom Status exportiert."""
        from app.models.invoice import Invoice, InvoiceStatus
        kunde = _datev_kunde(client)
        gs = _datev_rechnung(client, kunde, [("Preisnachlass", 1, "10.00", "REDUZIERT")],
                             typ="GUTSCHRIFT", finalisieren=False)
        with TestingSessionLocal() as db:
            db.get(Invoice, uuid.UUID(gs["id"])).status = InvoiceStatus.STORNIERT
            db.commit()

        _, _, zeilen = _datev_export(client)

        assert zeilen == []

    def test_proforma_und_entwurf_werden_nicht_exportiert(self, client):
        kunde = _datev_kunde(client)
        _datev_rechnung(client, kunde, DATEV_GEMISCHT, typ="PROFORMA")
        _datev_rechnung(client, kunde, DATEV_GEMISCHT, finalisieren=False)

        _, _, zeilen = _datev_export(client)

        assert zeilen == []

    def test_exportierte_rechnung_nur_mit_stornorechnung_stornierbar(self, client):
        """Nach dem Export kann nur ein Gegenbeleg den Umsatz in DATEV aufheben.
        Ein Storno ohne Stornorechnung nähme die Rechnung aus jedem weiteren
        Export, der Umsatz bliebe in DATEV still stehen."""
        kunde = _datev_kunde(client)
        rechnung = _datev_rechnung(client, kunde, DATEV_GEMISCHT)
        _datev_export(client)

        r = client.post(f"/api/v1/invoices/{rechnung['id']}/cancel",
                        json={"reason": "Doppelt erfasst", "create_credit_note": False})

        assert r.status_code == 400, r.text
        assert "DATEV" in r.json()["detail"]
        assert client.get(f"/api/v1/invoices/{rechnung['id']}").json()["status"] == "OFFEN"


class TestDatevWiederholungsexport:

    def _exportiert(self, rechnung_id):
        from app.models.invoice import Invoice
        with TestingSessionLocal() as db:
            inv = db.get(Invoice, uuid.UUID(rechnung_id))
            return inv.datev_exported, inv.datev_export_date

    def test_export_markiert_genau_die_exportierten(self, client):
        kunde = _datev_kunde(client)
        rechnung = _datev_rechnung(client, kunde, DATEV_GEMISCHT)
        entwurf = _datev_rechnung(client, kunde, DATEV_GEMISCHT, finalisieren=False)

        _datev_export(client)

        markiert, wann = self._exportiert(rechnung["id"])
        assert markiert is True and wann is not None
        assert self._exportiert(entwurf["id"]) == (False, None)

    def test_zweiter_lauf_ohne_wiederholung_ist_leer(self, client):
        kunde = _datev_kunde(client)
        _datev_rechnung(client, kunde, DATEV_GEMISCHT)
        _datev_export(client)

        daten, kopf, zeilen = _datev_export(client)

        assert kopf == DATEV_KOPF
        assert zeilen == []
        assert daten["record_count"] == 0

    def test_wiederholung_liefert_dieselben_zeilen_und_neue_dazu(self, client):
        kunde = _datev_kunde(client)
        _datev_rechnung(client, kunde, DATEV_GEMISCHT)
        _, _, erster_lauf = _datev_export(client)
        neu = _datev_rechnung(client, kunde, [("Kresse", 1, "10.00", "REDUZIERT")])

        _, _, nur_neu = _datev_export(client)
        _, _, alles = _datev_export(client, erneut=True)

        assert [z[9] for z in nur_neu] == [neu["invoice_number"]]
        assert alles == erster_lauf + nur_neu

    def test_zahlungen_werden_mit_wiederholt(self, client):
        kunde = _datev_kunde(client)
        rechnung = _datev_rechnung(client, kunde, DATEV_GEMISCHT)
        r = client.post(f"/api/v1/invoices/{rechnung['id']}/payments", json={
            "invoice_id": rechnung["id"], "amount": "33.89",
            "payment_date": date.today().isoformat(),
        })
        assert r.status_code == 201, r.text
        nr = rechnung["invoice_number"]
        zahlung = _datev_zeile("33,89", "S", "1200", "10008", nr, "Zahlung Ökoring Testkunde")

        _, _, erster = _datev_export(client, zahlungen=True)
        _, _, zweiter = _datev_export(client, zahlungen=True)
        _, _, wiederholt = _datev_export(client, zahlungen=True, erneut=True)

        assert erster[-1] == zahlung
        assert zweiter == []
        assert wiederholt == erster


# ---------------------------------------------------------------------------
# S6 — Doppelte Abrechnung verhindern
# ---------------------------------------------------------------------------

S6_PREVIEW = "/api/v1/invoices/batch-run/preview"
S6_COMMIT = "/api/v1/invoices/batch-run/commit"


def _s6_kunde(client, name="Ökoring Handels GmbH"):
    r = client.post("/api/v1/sales/customers", json={"name": name, "typ": "HANDEL"})
    assert r.status_code in (200, 201), r.text
    return r.json()


def _s6_bestellung(client, kunde, liefertag="2026-03-05", menge=10, preis=2.50):
    """Freitext-Position mit ausdrücklichem Satz — bleibt vom A3-Fix unberührt."""
    r = client.post("/api/v1/sales/orders", json={
        "customer_id": kunde["id"],
        "requested_delivery_date": liefertag,
        "lines": [{"product_name": "Erbsen-Schale", "quantity": menge, "unit": "STK",
                   "unit_price": preis, "tax_rate": "REDUZIERT"}],
    })
    assert r.status_code == 201, r.text
    return r.json()


def _s6_lieferschein(client, bestellung, zusaetzlich=False):
    params = {"zusaetzlich": "true"} if zusaetzlich else None
    r = client.post(f"/api/v1/sales/orders/{bestellung['id']}/delivery-notes",
                    json={}, params=params)
    assert r.status_code == 201, r.text
    return r.json()


def _s6_aus_bestellung(client, bestellung):
    return client.post(f"/api/v1/invoices/from-order/{bestellung['id']}")


def _s6_finalisieren(client, rechnung):
    r = client.post(f"/api/v1/invoices/{rechnung['id']}/finalize")
    assert r.status_code == 200, r.text
    return r.json()


def _s6_storno(client, rechnung):
    r = client.post(f"/api/v1/invoices/{rechnung['id']}/cancel", json={
        "reason": "Pfand mit falschem Steuersatz", "reason_code": "PREISFEHLER",
    })
    assert r.status_code == 200, r.text
    return r.json()


def _s6_lauf(client, endpoint):
    r = client.post(endpoint, json={"period_from": "2026-03-01", "period_to": "2026-03-31"})
    assert r.status_code in (200, 201), r.text
    return r.json()


def _s6_rechnung_ohne_bestellung(client, kunde):
    """Manuelle Rechnung ohne Bestellbezug — schnell, ohne Bestellanlage."""
    r = client.post("/api/v1/invoices", json={
        "customer_id": kunde["id"], "invoice_date": date.today().isoformat(),
    })
    assert r.status_code == 201, r.text
    return r.json()


def _s6_pdf_text(client, rechnung):
    from tests.test_documents_preise import _pdf_text
    r = client.get(f"/api/v1/invoices/{rechnung['id']}/pdf")
    assert r.status_code == 200, r.text
    return _pdf_text(r.content)


def _s6_ls_an_rechnung(client, rechnung):
    r = client.get(f"/api/v1/invoices/{rechnung['id']}/delivery-notes")
    assert r.status_code == 200, r.text
    return sorted(n["delivery_note_number"] for n in r.json())


class TestS6KeineZweiteRechnungZurBestellung:
    """(a) from-order lehnt ab, solange eine nicht stornierte Rechnung existiert."""

    def test_zweite_rechnung_aus_bestellung_wird_abgelehnt(self, client):
        bestellung = _s6_bestellung(client, _s6_kunde(client))
        erste = _s6_aus_bestellung(client, bestellung)
        assert erste.status_code == 201, erste.text

        zweite = _s6_aus_bestellung(client, bestellung)

        assert zweite.status_code == 409, zweite.text
        assert erste.json()["invoice_number"] in zweite.json()["detail"]
        alle = client.get("/api/v1/invoices",
                          params={"customer_id": bestellung["customer_id"]}).json()
        assert len(alle) == 1, "Die abgelehnte Rechnung darf nicht gespeichert sein"

    def test_finalisierte_rechnung_sperrt_ebenso(self, client):
        bestellung = _s6_bestellung(client, _s6_kunde(client))
        _s6_finalisieren(client, _s6_aus_bestellung(client, bestellung).json())

        assert _s6_aus_bestellung(client, bestellung).status_code == 409

    def test_nach_storno_ist_neuausstellung_erlaubt(self, client):
        """Runbook S3: RE-00002/4 stornieren und neu ausstellen."""
        bestellung = _s6_bestellung(client, _s6_kunde(client))
        erste = _s6_finalisieren(client, _s6_aus_bestellung(client, bestellung).json())
        _s6_storno(client, erste)

        neu = _s6_aus_bestellung(client, bestellung)

        assert neu.status_code == 201, neu.text
        assert neu.json()["id"] != erste["id"]
        # ... und die Neuausstellung sperrt wieder
        assert _s6_aus_bestellung(client, bestellung).status_code == 409

    def test_sammelrechnung_sperrt_rechnung_aus_bestellung(self, client):
        bestellung = _s6_bestellung(client, _s6_kunde(client))
        _s6_lieferschein(client, bestellung)
        sammel = _s6_lauf(client, S6_COMMIT)["rechnungen"][0]

        r = _s6_aus_bestellung(client, bestellung)

        assert r.status_code == 409, r.text
        assert sammel["invoice_number"] in r.json()["detail"]

    def test_manuelle_rechnung_mit_bestellbezug_wird_abgelehnt(self, client):
        bestellung = _s6_bestellung(client, _s6_kunde(client))
        assert _s6_aus_bestellung(client, bestellung).status_code == 201

        r = client.post("/api/v1/invoices", json={
            "customer_id": bestellung["customer_id"], "order_id": bestellung["id"],
            "invoice_date": date.today().isoformat(),
        })

        assert r.status_code == 409, r.text

    def test_bestellstatus_bleibt_und_lieferung_ist_weiter_moeglich(self, client):
        """(f) Kein automatisches FAKTURIERT: Rechnung vor Lieferung darf das
        Quittieren (GELIEFERT + Bestandsabbuchung) nicht aushebeln."""
        bestellung = _s6_bestellung(client, _s6_kunde(client),
                                    liefertag=date.today().isoformat())
        r = client.post(f"/api/v1/sales/orders/{bestellung['id']}/confirm")
        assert r.status_code == 200, r.text
        ls = _s6_lieferschein(client, bestellung)
        assert _s6_aus_bestellung(client, bestellung).status_code == 201

        assert client.get(f"/api/v1/sales/orders/{bestellung['id']}").json()["status"] == "BESTAETIGT"

        r = client.patch(f"/api/v1/sales/delivery-notes/{ls['id']}/mark-delivered",
                         json={"signed_by": "Fahrer"})
        assert r.status_code == 200, r.text
        assert client.get(f"/api/v1/sales/orders/{bestellung['id']}").json()["status"] == "GELIEFERT"


class TestS6LieferscheinHaengtAnDerRechnung:
    """(b) from-order verknüpft den Lieferschein; Storno löst, Neuausstellung verknüpft neu."""

    def test_rechnung_aus_bestellung_verknuepft_den_lieferschein(self, client):
        bestellung = _s6_bestellung(client, _s6_kunde(client))
        ls = _s6_lieferschein(client, bestellung)

        rechnung = _s6_aus_bestellung(client, bestellung).json()

        assert _s6_ls_an_rechnung(client, rechnung) == [ls["delivery_note_number"]]

    def test_storno_loest_und_neuausstellung_verknuepft_neu(self, client):
        bestellung = _s6_bestellung(client, _s6_kunde(client))
        ls = _s6_lieferschein(client, bestellung)
        erste = _s6_finalisieren(client, _s6_aus_bestellung(client, bestellung).json())

        _s6_storno(client, erste)
        assert _s6_ls_an_rechnung(client, erste) == []

        neu = _s6_aus_bestellung(client, bestellung).json()
        assert _s6_ls_an_rechnung(client, neu) == [ls["delivery_note_number"]]

    def test_bei_zwei_lieferscheinen_haengt_genau_einer_an(self, client):
        """Die PDF-Tabelle 'Enthaltene Lieferscheine' zeigt je Lieferschein die
        volle Bestellsumme — zwei Einträge würden sie doppelt ausweisen."""
        bestellung = _s6_bestellung(client, _s6_kunde(client))
        ls1 = _s6_lieferschein(client, bestellung)
        _s6_lieferschein(client, bestellung, zusaetzlich=True)

        rechnung = _s6_aus_bestellung(client, bestellung).json()

        assert _s6_ls_an_rechnung(client, rechnung) == [ls1["delivery_note_number"]]

    def test_pdf_mit_bestellbezug_bekommt_keine_lieferscheinanlage(self, client):
        """GoBD: die Anlage liest DeliveryNote.invoice_id, und die löst ein
        Storno wieder. Bei Rechnungen mit Bestellbezug würde sich das PDF
        nachträglich ändern — deshalb dort keine Anlage."""
        bestellung = _s6_bestellung(client, _s6_kunde(client))
        _s6_lieferschein(client, bestellung)
        rechnung = _s6_aus_bestellung(client, bestellung).json()

        assert b"Enthaltene Lieferscheine" not in _s6_pdf_text(client, rechnung)

    def test_sammelrechnung_behaelt_die_lieferscheinanlage(self, client):
        bestellung = _s6_bestellung(client, _s6_kunde(client))
        _s6_lieferschein(client, bestellung)
        sammel = _s6_lauf(client, S6_COMMIT)["rechnungen"][0]

        assert b"Enthaltene Lieferscheine" in _s6_pdf_text(client, sammel)


class TestS6SammellaufRechnetJedeBestellungEinmal:
    """(c) Vorschau und Festschreiben: abgerechnete Bestellungen raus, je Bestellung einmal."""

    def test_bestellung_mit_rechnung_aus_bestellung_wird_uebersprungen(self, client):
        bestellung = _s6_bestellung(client, _s6_kunde(client))
        _s6_lieferschein(client, bestellung)
        assert _s6_aus_bestellung(client, bestellung).status_code == 201

        assert _s6_lauf(client, S6_PREVIEW)["kunden"] == []
        assert _s6_lauf(client, S6_COMMIT)["rechnungen"] == []

    def test_altbestand_ohne_verknuepfung_wird_uebersprungen(self, client):
        """RE-00002/3/4: vor S6 aus Bestellung erzeugt, Lieferschein ohne invoice_id."""
        from uuid import UUID
        from app.models.documents import DeliveryNote
        from tests.conftest import TestingSessionLocal

        bestellung = _s6_bestellung(client, _s6_kunde(client))
        ls = _s6_lieferschein(client, bestellung)
        assert _s6_aus_bestellung(client, bestellung).status_code == 201
        with TestingSessionLocal() as db:
            db.get(DeliveryNote, UUID(ls["id"])).invoice_id = None
            db.commit()

        assert _s6_lauf(client, S6_PREVIEW)["kunden"] == []
        assert _s6_lauf(client, S6_COMMIT)["rechnungen"] == []

    def test_nach_storno_rechnet_der_lauf_die_bestellung_wieder(self, client):
        bestellung = _s6_bestellung(client, _s6_kunde(client))
        _s6_lieferschein(client, bestellung)
        _s6_storno(client, _s6_finalisieren(client, _s6_aus_bestellung(client, bestellung).json()))

        kunden = _s6_lauf(client, S6_PREVIEW)["kunden"]

        assert len(kunden) == 1
        assert kunden[0]["anzahl_lieferscheine"] == 1

    def test_zwei_lieferscheine_einer_bestellung_werden_einmal_berechnet(self, client):
        bestellung = _s6_bestellung(client, _s6_kunde(client), menge=10, preis=2.50)
        ls1 = _s6_lieferschein(client, bestellung)
        _s6_lieferschein(client, bestellung, zusaetzlich=True)

        k = _s6_lauf(client, S6_PREVIEW)["kunden"][0]
        assert k["anzahl_lieferscheine"] == 1
        assert Decimal(str(k["summe_netto"])) == Decimal("25.00")

        rechnung = _s6_lauf(client, S6_COMMIT)["rechnungen"][0]
        assert Decimal(str(rechnung["subtotal"])) == Decimal("25.00")
        assert _s6_ls_an_rechnung(client, rechnung) == [ls1["delivery_note_number"]]
        assert _s6_lauf(client, S6_COMMIT)["rechnungen"] == []

    def test_quittierter_lieferschein_vertritt_die_bestellung(self, client):
        bestellung = _s6_bestellung(client, _s6_kunde(client))
        # Manager 08.10. (Paket 2, Entscheidung E5): Quittieren setzt ab Paket 2 eine
        # bestätigte Bestellung voraus — ENTWURF → GELIEFERT ist kein erlaubter Übergang.
        r = client.post(f"/api/v1/sales/orders/{bestellung['id']}/confirm")
        assert r.status_code == 200, r.text
        _s6_lieferschein(client, bestellung)
        ls2 = _s6_lieferschein(client, bestellung, zusaetzlich=True)
        r = client.patch(f"/api/v1/sales/delivery-notes/{ls2['id']}/mark-delivered",
                         json={"signed_by": "Fahrer", "actual_delivery_date": "2026-03-06"})
        assert r.status_code == 200, r.text

        rechnung = _s6_lauf(client, S6_COMMIT)["rechnungen"][0]

        assert _s6_ls_an_rechnung(client, rechnung) == [ls2["delivery_note_number"]]

    def test_fakturierte_bestellung_ohne_rechnung_wird_uebersprungen(self, client):
        """Spec-Nachtrag 08.10.2026: FAKTURIERT heißt abgerechnet — auch ohne
        Rechnung im System (z. B. außerhalb von NovaERP berechnet). Bisher
        rechnete der Lauf sie trotzdem ab, sobald sie einen freien Lieferschein
        hatte (Nachtrag T4, Probe P7)."""
        from uuid import UUID
        from app.models.enums import OrderStatus
        from app.models.order import Order

        bestellung = _s6_bestellung(client, _s6_kunde(client))
        _s6_lieferschein(client, bestellung)
        # Zustand wie nach dem manuellen Übergang GELIEFERT → FAKTURIERT
        with TestingSessionLocal() as db:
            db.get(Order, UUID(bestellung["id"])).status = OrderStatus.FAKTURIERT
            db.commit()

        assert _s6_lauf(client, S6_PREVIEW)["kunden"] == []
        assert _s6_lauf(client, S6_COMMIT)["rechnungen"] == []


class TestS6ZweiterLieferschein:
    """(d) Ein weiterer Lieferschein nur mit ausdrücklicher Bestätigung."""

    def test_zweiter_lieferschein_braucht_bestaetigung(self, client):
        bestellung = _s6_bestellung(client, _s6_kunde(client))
        ls1 = _s6_lieferschein(client, bestellung)

        r = client.post(f"/api/v1/sales/orders/{bestellung['id']}/delivery-notes", json={})

        assert r.status_code == 409, r.text
        assert ls1["delivery_note_number"] in r.json()["detail"]
        alle = client.get(f"/api/v1/sales/orders/{bestellung['id']}/delivery-notes").json()
        assert len(alle) == 1

    def test_mit_bestaetigung_wird_er_angelegt(self, client):
        bestellung = _s6_bestellung(client, _s6_kunde(client))
        ls1 = _s6_lieferschein(client, bestellung)

        ls2 = _s6_lieferschein(client, bestellung, zusaetzlich=True)

        assert ls2["delivery_note_number"] != ls1["delivery_note_number"]


class TestS6RechnungenJeBestellung:
    """(e) GET /invoices?order_id= — Grundlage des Belege-Dialogs."""

    def test_filter_findet_die_rechnung_jenseits_der_ersten_seite(self, client):
        """Der Dialog filterte die 20 neuesten Rechnungen — ab Rechnung 21 fehlte sie."""
        kunde = _s6_kunde(client)
        ziel = _s6_bestellung(client, kunde)
        rechnung = _s6_aus_bestellung(client, ziel).json()
        for _ in range(21):
            _s6_rechnung_ohne_bestellung(client, kunde)

        r = client.get("/api/v1/invoices", params={"order_id": ziel["id"]})

        assert r.status_code == 200, r.text
        assert [i["id"] for i in r.json()] == [rechnung["id"]]

    def test_filter_findet_die_sammelrechnung_ueber_den_lieferschein(self, client):
        bestellung = _s6_bestellung(client, _s6_kunde(client))
        _s6_lieferschein(client, bestellung)
        sammel = _s6_lauf(client, S6_COMMIT)["rechnungen"][0]
        _s6_rechnung_ohne_bestellung(client, _s6_kunde(client, name="Großer Kern"))

        r = client.get("/api/v1/invoices", params={"order_id": bestellung["id"]})

        assert [i["id"] for i in r.json()] == [sammel["id"]]

    def test_stornierte_und_neue_rechnung_erscheinen_beide(self, client):
        bestellung = _s6_bestellung(client, _s6_kunde(client))
        erste = _s6_finalisieren(client, _s6_aus_bestellung(client, bestellung).json())
        _s6_storno(client, erste)
        neu = _s6_aus_bestellung(client, bestellung).json()
        _s6_rechnung_ohne_bestellung(client, _s6_kunde(client, name="Großer Kern"))

        r = client.get("/api/v1/invoices", params={"order_id": bestellung["id"]})

        rechnungen = {(i["id"], i["status"]) for i in r.json() if i["invoice_type"] == "RECHNUNG"}
        assert rechnungen == {(erste["id"], "STORNIERT"), (neu["id"], "ENTWURF")}
