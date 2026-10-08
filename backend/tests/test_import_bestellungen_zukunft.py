import io
from datetime import datetime, timedelta
from decimal import Decimal
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from tests.conftest import TestingSessionLocal


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _heute_berlin():
    return datetime.now(ZoneInfo("Europe/Berlin")).date()


def _kunde(client, name="Testküche Zukunft"):
    response = client.post("/api/v1/sales/customers", json={"name": name, "typ": "GASTRO"})
    assert response.status_code == 201, response.text
    return response.json()


def _produkt(sku="MG-ZUKUNFT"):
    from app.models.product import Product, ProductCategory
    from app.models.unit import UnitCategory, UnitOfMeasure

    with TestingSessionLocal() as db:
        unit = UnitOfMeasure(
            code=f"STK-{sku}",
            name=f"Stück {sku}",
            symbol="Stk",
            category=UnitCategory.COUNT,
            is_base_unit=True,
        )
        db.add(unit)
        db.flush()
        product = Product(
            sku=sku,
            name=f"Produkt {sku}",
            category=ProductCategory.MICROGREEN,
            base_unit_id=unit.id,
            base_price=Decimal("3.90"),
        )
        db.add(product)
        db.commit()
        db.refresh(product)
        return {"id": str(product.id), "sku": product.sku, "name": product.name}


def _import_bestellungen(client, rows):
    template = client.get("/api/v1/imports/template/order_history")
    assert template.status_code == 200, template.text
    workbook = load_workbook(io.BytesIO(template.content))
    sheet = workbook["Daten"]
    for row in rows:
        sheet.append(row)
    content = io.BytesIO()
    workbook.save(content)
    return client.post(
        "/api/v1/imports/order_history",
        files={"file": ("bestellungen.xlsx", content.getvalue(), XLSX_MIME)},
    )


def _zeile(externe_nummer, kunde, produkt, bestelldatum, lieferdatum, status=None):
    return [
        externe_nummer,
        kunde["name"],
        bestelldatum,
        lieferdatum,
        produkt["sku"],
        2,
        "STK",
        "3.90",
        status,
        None,
    ]


def _bestellung(client, externe_nummer):
    response = client.get("/api/v1/sales/orders")
    assert response.status_code == 200, response.text
    summary = next(item for item in response.json()["items"] if item["customer_reference"] == externe_nummer)
    response = client.get(f"/api/v1/sales/orders/{summary['id']}")
    assert response.status_code == 200, response.text
    return response.json()


def test_leerer_status_zukunft_wird_bestaetigt_und_bleibt_bearbeitbar(client):
    heute = _heute_berlin()
    lieferdatum = heute + timedelta(days=3)
    kunde = _kunde(client)
    produkt = _produkt()
    externe_nummer = "ZUK-LEER-001"

    response = _import_bestellungen(
        client,
        [_zeile(externe_nummer, kunde, produkt, heute, lieferdatum)],
    )
    assert response.status_code == 200, response.text

    order = _bestellung(client, externe_nummer)
    assert order["status"] == "BESTAETIGT"
    assert order["actual_delivery_date"] is None
    assert order["confirmed_delivery_date"] == lieferdatum.isoformat()

    plan_date = lieferdatum - timedelta(days=1)
    response = client.get("/api/v1/production/day-plan", params={"target_date": plan_date.isoformat()})
    assert response.status_code == 200, response.text
    assert any(item["order_id"] == order["id"] for item in response.json()["verpacken"])

    response = client.post(
        f"/api/v1/sales/orders/{order['id']}/lines",
        json={
            "product_id": produkt["id"],
            "product_name": produkt["name"],
            "quantity": 1,
            "unit": "STK",
            "unit_price": "3.90",
        },
    )
    assert response.status_code == 201, response.text


def test_leerer_status_vergangenheit_wird_geliefert(client):
    heute = _heute_berlin()
    lieferdatum = heute - timedelta(days=3)
    kunde = _kunde(client)
    produkt = _produkt()
    externe_nummer = "ALT-LEER-001"

    response = _import_bestellungen(
        client,
        [_zeile(externe_nummer, kunde, produkt, lieferdatum, lieferdatum)],
    )
    assert response.status_code == 200, response.text

    order = _bestellung(client, externe_nummer)
    assert order["status"] == "GELIEFERT"
    assert order["actual_delivery_date"] == lieferdatum.isoformat()


def test_leerer_status_heute_wird_bestaetigt(client):
    heute = _heute_berlin()
    kunde = _kunde(client)
    produkt = _produkt()
    externe_nummer = "HEUTE-LEER-001"

    response = _import_bestellungen(
        client,
        [_zeile(externe_nummer, kunde, produkt, heute, heute)],
    )
    assert response.status_code == 200, response.text

    order = _bestellung(client, externe_nummer)
    assert order["status"] == "BESTAETIGT"
    assert order["actual_delivery_date"] is None
    assert order["confirmed_delivery_date"] == heute.isoformat()


def test_geliefert_mit_zukunftsdatum_lehnt_ganzen_lauf_ab(client):
    heute = _heute_berlin()
    lieferdatum = heute + timedelta(days=3)
    kunde = _kunde(client)
    produkt = _produkt()
    gueltige_nummer = "ROLLBACK-OK-001"
    fehlerhafte_nummer = "ROLLBACK-FEHLER-001"

    response = _import_bestellungen(
        client,
        [
            _zeile(gueltige_nummer, kunde, produkt, heute, lieferdatum, "BESTAETIGT"),
            _zeile(fehlerhafte_nummer, kunde, produkt, heute, lieferdatum, "GELIEFERT"),
        ],
    )

    assert response.status_code == 400, response.text
    assert fehlerhafte_nummer in response.json()["detail"]
    assert lieferdatum.strftime("%d.%m.%Y") in response.json()["detail"]
    orders = client.get("/api/v1/sales/orders").json()["items"]
    references = {item["customer_reference"] for item in orders}
    assert gueltige_nummer not in references
    assert fehlerhafte_nummer not in references


def test_ausdruecklich_bestaetigt_in_vergangenheit_bleibt_bestaetigt(client):
    heute = _heute_berlin()
    lieferdatum = heute - timedelta(days=3)
    kunde = _kunde(client)
    produkt = _produkt()
    externe_nummer = "ALT-BESTAETIGT-001"

    response = _import_bestellungen(
        client,
        [_zeile(externe_nummer, kunde, produkt, lieferdatum, lieferdatum, "BESTAETIGT")],
    )
    assert response.status_code == 200, response.text

    order = _bestellung(client, externe_nummer)
    assert order["status"] == "BESTAETIGT"
    assert order["actual_delivery_date"] is None


def test_vorlage_enthaelt_leeren_status_mit_zukunftsdatum(client):
    heute = _heute_berlin()
    response = client.get("/api/v1/imports/template/order_history")
    assert response.status_code == 200, response.text

    workbook = load_workbook(io.BytesIO(response.content), data_only=True)
    sheet = workbook["Beispiel"]
    headers = [str(cell.value).split(" ")[0] for cell in sheet[1]]
    lieferdatum_index = headers.index("lieferdatum")
    status_index = headers.index("status")
    beispielzeilen = list(sheet.iter_rows(min_row=2, max_row=4, values_only=True))

    assert any(
        row[status_index] in (None, "")
        and datetime.strptime(row[lieferdatum_index], "%d.%m.%Y").date() > heute
        for row in beispielzeilen
    )
